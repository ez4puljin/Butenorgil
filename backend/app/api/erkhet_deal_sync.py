"""Erkhet sync шалгалт — erxes deal → Эрхэт авлага (app/services/erkhet_deal_check).

  GET  /erkhet-sync/meta               — board/pipeline/stage (ангилалтай), ebarimtConfig (нууцгүй), dateType
  PUT  /erkhet-sync/settings           — stage-ийн ангилал, pipeline хасах/нэмэх (admin)
  POST /erkhet-sync/check              — {boardId, pipelineIds?, stageIds?, dateField, startDate, endDate} → run_id
  GET  /erkhet-sync/runs               — сүүлийн шалгалтууд
  GET  /erkhet-sync/runs/{id}          — явц + (дууссан бол) deal бүрийн үр дүн
  GET  /erkhet-sync/runs/{id}/xlsx     — Excel
  POST /erkhet-sync/resync/preview     — dry-run (erxes-т юу ч бичихгүй)
  POST /erkhet-sync/resync             — баталгаажуулсны дараа дахин sync (admin) → run_id
  GET  /erkhet-sync/actions            — дахин sync-ийн audit

Хатуу дүрэм: автоматаар sync ХИЙХГҮЙ — зөвхөн хэрэглэгч сонгож dry-run баталсны дараа (confirm=true).
"""
from __future__ import annotations

import io
import json
import threading
from datetime import date, datetime, timedelta
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.api.deps import get_db, require_role
from app.core.audit import audit
from app.models.erkhet_sync_check import ErkhetSyncAction, ErkhetSyncItem, ErkhetSyncRun, ErkhetSyncSetting
from app.models.user import User
from app.services import erkhet_deal_check as svc

router = APIRouter(prefix="/erkhet-sync", tags=["erkhet-sync"])
VIEW = ("admin", "supervisor", "manager")

_meta_cache: dict = {"at": None, "val": None}


def _meta(fresh: bool = False) -> dict:
    now = datetime.utcnow()
    if fresh or not _meta_cache["val"] or (now - _meta_cache["at"]).total_seconds() > 600:
        _meta_cache.update(at=now, val=svc.fetch_meta())
    return _meta_cache["val"]


@router.get("/meta")
def meta(fresh: bool = False, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    try:
        m = _meta(fresh)
    except Exception as e:
        raise HTTPException(502, f"erxes-тэй холбогдож чадсангүй: {e}")
    saved_st, saved_pl = svc.load_settings(db)
    boards = []
    for b in m["boards"]:
        pipes = []
        for p in b["pipelines"]:
            inc = saved_pl.get(p["_id"], "include" if svc.default_pipeline_included(p["name"]) else "exclude")
            pipes.append({"_id": p["_id"], "name": p["name"], "included": inc == "include",
                          "stages": [{"_id": s["_id"], "name": s["name"],
                                      "category": svc.stage_category(s["_id"], s["name"], saved_st),
                                      "saved": s["_id"] in saved_st} for s in p["stages"]]})
        boards.append({"_id": b["_id"], "name": b["name"], "pipelines": pipes})
    default_board = next((b["_id"] for b in boards if svc.BOARD_DEFAULT in b["name"]), boards[0]["_id"] if boards else "")
    return {"boards": boards, "defaultBoardId": default_board, "configs": m["configs"],
            "dateTypes": [{"value": v, "label": l} for v, l in svc.DATE_TYPES],
            "statuses": {k: {"severity": s, "label": l} for k, (s, l) in svc.STATUS_INFO.items()},
            "canResync": (getattr(u, "base_role", None) or u.role) == "admin"}


class SettingIn(BaseModel):
    kind: str            # stage | pipeline
    ref_id: str
    name: str = ""
    value: str           # stage: expected|pending|cancel; pipeline: include|exclude


@router.put("/settings")
def save_settings(body: list[SettingIn], db: Session = Depends(get_db), u: User = Depends(require_role("admin"))):
    for s in body:
        ok = (s.kind == "stage" and s.value in svc.CATEGORIES) or (s.kind == "pipeline" and s.value in ("include", "exclude"))
        if not ok:
            raise HTTPException(400, f"Буруу тохиргоо: {s.kind}={s.value}")
        r = db.query(ErkhetSyncSetting).filter(ErkhetSyncSetting.kind == s.kind, ErkhetSyncSetting.ref_id == s.ref_id).first()
        if r is None:
            r = ErkhetSyncSetting(kind=s.kind, ref_id=s.ref_id)
            db.add(r)
        r.name, r.value, r.updated_by, r.updated_at = s.name[:300], s.value, u.username, datetime.utcnow()
    db.commit()
    return {"ok": True, "saved": len(body)}


class CheckIn(BaseModel):
    boardId: str = ""
    pipelineIds: list[str] = []
    stageIds: list[str] = []
    dateField: str = "createdAt"
    startDate: str
    endDate: str


def _running(db: Session, kind: str) -> Optional[ErkhetSyncRun]:
    since = datetime.utcnow() - timedelta(hours=2)
    return db.query(ErkhetSyncRun).filter(ErkhetSyncRun.kind == kind, ErkhetSyncRun.status == "running",
                                          ErkhetSyncRun.created_at >= since).first()


@router.post("/check")
def start_check(body: CheckIn, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    if body.dateField not in svc.DATE_FIELDS:
        raise HTTPException(400, "Огнооны талбар буруу.")
    try:
        a, b = date.fromisoformat(body.startDate), date.fromisoformat(body.endDate)
    except ValueError:
        raise HTTPException(400, "Огноо буруу (YYYY-MM-DD).")
    if a > b or (b - a).days > 62:
        raise HTTPException(400, "Огнооны хүрээ 1–62 хоног байна.")
    busy = _running(db, "check")
    if busy:
        return {"run_id": busy.id, "reused": True}
    r = ErkhetSyncRun(kind="check", status="running", params=body.model_dump_json(), created_by=u.username)
    db.add(r)
    db.commit()
    threading.Thread(target=svc.run_check, args=(r.id, body.model_dump()), daemon=True,
                     name=f"erkhet-check-{r.id}").start()
    return {"run_id": r.id, "reused": False}


def _run_out(r: ErkhetSyncRun) -> dict:
    return {"id": r.id, "kind": r.kind, "status": r.status, "params": json.loads(r.params or "{}"),
            "progress": r.progress, "done_n": r.done_n, "total_n": r.total_n,
            "summary": json.loads(r.summary or "{}"), "error": r.error, "created_by": r.created_by,
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "finished_at": r.finished_at.isoformat() if r.finished_at else None}


@router.get("/runs")
def list_runs(kind: str = "check", limit: int = 20, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    rs = db.query(ErkhetSyncRun).filter(ErkhetSyncRun.kind == kind).order_by(ErkhetSyncRun.id.desc()).limit(min(limit, 100))
    return [_run_out(r) for r in rs]


@router.get("/runs/{run_id}")
def get_run(run_id: int, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    r = db.get(ErkhetSyncRun, run_id)
    if not r:
        raise HTTPException(404, "Олдсонгүй.")
    out = _run_out(r)
    if r.status == "done" and r.kind == "check":
        out["items"] = [_item_out(x) for x in db.query(ErkhetSyncItem).filter(ErkhetSyncItem.run_id == run_id)
                        .order_by(ErkhetSyncItem.id)]
    if r.kind == "resync":
        out["actions"] = [_action_out(a) for a in db.query(ErkhetSyncAction).filter(ErkhetSyncAction.run_id == run_id)]
    return out


def _item_out(x: ErkhetSyncItem) -> dict:
    return {"deal_id": x.deal_id, "number": x.number, "name": x.name, "pipeline_id": x.pipeline_id,
            "stage_id": x.stage_id, "category": x.category, "amount": x.amount, "severity": x.severity,
            "statuses": json.loads(x.statuses or "[]"), "data": json.loads(x.data or "{}")}


def _action_out(a: ErkhetSyncAction) -> dict:
    return {"deal_id": a.deal_id, "number": a.number, "config_stage_id": a.config_stage_id,
            "config_title": a.config_title, "date_type": a.date_type, "payments_before": a.payments_before,
            "payments_after": a.payments_after, "before": json.loads(a.before_statuses or "[]"),
            "after": json.loads(a.after_statuses or "[]"), "result": a.result, "message": a.message,
            "user": a.user_name, "at": a.created_at.isoformat() if a.created_at else None}


@router.get("/runs/{run_id}/xlsx")
def export_xlsx(run_id: int, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    r = db.get(ErkhetSyncRun, run_id)
    if not r or r.status != "done":
        raise HTTPException(404, "Дууссан шалгалт олдсонгүй.")
    p = json.loads(r.params or "{}")
    wb = Workbook()
    ws = wb.active
    ws.title = "Erkhet sync"
    head = ["Deal №", "Нэр", "Pipeline", "Stage", "Ангилал", "Дүн", "Статус", "Ноцтой", "Erkhet bill №",
            "Erkhet харилцагч", "Erkhet огноо", "Сүүлийн sync", "Илгээсэн төлбөр", "Erkhet хариу",
            "Үүссэн (УБ)", "Тайлбар"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    fills = {"error": "FDE2E2", "warn": "FEF3C7", "info": "E0F2FE"}
    for x in db.query(ErkhetSyncItem).filter(ErkhetSyncItem.run_id == run_id).order_by(ErkhetSyncItem.id):
        d = json.loads(x.data or "{}")
        sts = json.loads(x.statuses or "[]")
        ws.append([x.number, x.name, d.get("pipeline", ""), d.get("stage", ""), x.category, x.amount,
                   ", ".join(svc.STATUS_INFO.get(s, ("", s))[1] for s in sts), x.severity, d.get("erkhet_bill", ""),
                   d.get("erkhet_customer", ""), d.get("erkhet_date", ""), svc.ub_date(d.get("last_sync_at")),
                   d.get("pay_label", ""), ", ".join(d.get("resp_codes") or []), svc.ub_date(d.get("createdAt")),
                   "; ".join(d.get("notes") or [])])
        ws.cell(ws.max_row, 6).number_format = "#,##0"
        if x.severity in fills:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", start_color=fills[x.severity])
    for col, w in zip("ABCDEFGHIJKLMNOP", (16, 34, 22, 22, 10, 13, 36, 8, 16, 12, 12, 12, 14, 14, 12, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"Erkhet_sync_{p.get('startDate', '')}_{p.get('endDate', '')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fn)}"})


class ResyncIn(BaseModel):
    run_id: int                        # аль шалгалтаас (deal-ийн pipeline, статус)
    deal_ids: list[str]
    configs: dict[str, str]            # pipelineId → configStageId
    dateType: str = "created"
    fixPayments: bool = True
    confirm: bool = False


def _resync_items(db: Session, body: ResyncIn) -> list[dict]:
    ids = list(dict.fromkeys(body.deal_ids))
    if not ids:
        raise HTTPException(400, "Deal сонгоогүй байна.")
    if len(ids) > 300:
        raise HTTPException(400, "Нэг удаад 300-аас ихгүй deal.")
    rows = {x.deal_id: x for x in db.query(ErkhetSyncItem).filter(ErkhetSyncItem.run_id == body.run_id,
                                                                    ErkhetSyncItem.deal_id.in_(ids))}
    missing = [i for i in ids if i not in rows]
    if missing:
        raise HTTPException(400, f"{len(missing)} deal энэ шалгалтад алга.")
    return [{"deal_id": i, "number": rows[i].number, "pipeline_id": rows[i].pipeline_id, "amount": rows[i].amount,
             "statuses": json.loads(rows[i].statuses or "[]"),
             "payments_bad": json.loads(rows[i].data or "{}").get("payments_bad", False)} for i in ids]


@router.post("/resync/preview")
def resync_preview(body: ResyncIn, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    """Dry-run: erxes-т юу ч бичихгүй."""
    if body.dateType not in {v for v, _ in svc.DATE_TYPES}:
        raise HTTPException(400, "dateType буруу.")
    items = _resync_items(db, body)
    cfgs = {c["stageId"]: c for c in _meta()["configs"]}
    groups: dict[str, dict] = {}
    no_cfg = []
    for x in items:
        cid = body.configs.get(x["pipeline_id"], "")
        c = cfgs.get(cid)
        if not c:
            no_cfg.append(x["number"])
            continue
        g = groups.setdefault(cid, {"config_stage_id": cid, "title": c["title"], "defaultPay": c["defaultPay"],
                                    "cash_warning": c["defaultPay"] != "debtAmount", "n": 0, "amount": 0.0})
        g["n"] += 1
        g["amount"] += x["amount"]
    fix = [x["number"] for x in items if body.fixPayments and x["payments_bad"]]
    return {"n": len(items), "amount": round(sum(x["amount"] for x in items), 2), "dateType": body.dateType,
            "groups": list(groups.values()), "no_config": no_cfg, "fix_payments": fix,
            "batches": sum((g["n"] + svc.SYNC_BATCH - 1) // svc.SYNC_BATCH for g in groups.values())}


@router.post("/resync")
def resync(body: ResyncIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_role("admin"))):
    if not body.confirm:
        raise HTTPException(400, "Dry-run-ийг баталгаажуулна уу (confirm).")
    if body.dateType not in {v for v, _ in svc.DATE_TYPES}:
        raise HTTPException(400, "dateType буруу.")
    items = _resync_items(db, body)
    cfgs = {c["stageId"] for c in _meta()["configs"]}
    bad = [x["number"] for x in items if body.configs.get(x["pipeline_id"]) not in cfgs]
    if bad:
        raise HTTPException(400, f"Config сонгоогүй deal: {', '.join(bad[:5])}")
    busy = svc.lock_deals([x["deal_id"] for x in items])
    if busy:
        raise HTTPException(409, f"{len(busy)} deal дээр өөр sync ажиллаж байна.")
    params = {"items": items, "configs": body.configs, "dateType": body.dateType, "fixPayments": body.fixPayments,
              "check_run_id": body.run_id}
    r = ErkhetSyncRun(kind="resync", status="running", params=json.dumps(params, ensure_ascii=False),
                      total_n=len(items), created_by=u.username)
    db.add(r)
    db.commit()
    audit(db, request, u, action="erkhet_deal_resync", entity_type="erkhet_sync",
          extra={"run_id": r.id, "deals": len(items), "configs": body.configs, "dateType": body.dateType,
                 "fixPayments": body.fixPayments}, autocommit=True)
    threading.Thread(target=svc.run_resync, args=(r.id, params, u.username), daemon=True,
                     name=f"erkhet-resync-{r.id}").start()
    return {"run_id": r.id}


@router.get("/actions")
def actions(limit: int = 200, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    return [_action_out(a) for a in db.query(ErkhetSyncAction).order_by(ErkhetSyncAction.id.desc()).limit(min(limit, 1000))]
