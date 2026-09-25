"""Поддон хураалт — API.

  GET    /pallets/templates                 — загварууд
  POST   /pallets/templates                 — шинэ загвар
  PUT    /pallets/templates/{id}            — засах
  DELETE /pallets/templates/{id}            — устгах (ашиглагдаж байвал 400)
  GET    /pallets/lookup?q=&tag=            — баркод/код → бараа + одоогийн тохиргоо + дүн + сарын борлуулалт
  GET    /pallets/worklist?tag=&status=&q=  — tag-ийн бүх бараа борлуулалтын эрэмбээр, оруулсан/оруулаагүй
  GET    /pallets/products?q=&template_id=  — тохиргоотой бараануудын жагсаалт (дүнтэй)
  PUT    /pallets/products/{item_code}      — тохиргоо хадгалах (upsert)
  DELETE /pallets/products/{item_code}
  GET    /pallets/copy-candidates?source=&q=&scope=all|brand|match — хуулах бараа (бүх бараа, хуудаслалттай)
  POST   /pallets/products/{item_code}/copy — сонгосон талбаруудыг өөр бараанд хуулах
  GET    /pallets/export?tag=&status=       — Excel (tag-гүй бол тохиргоотой бүх бараа)
  GET    /pallets/sales-settings            — сарын дундажийн сарууд, тусгай сар, поддоноор тооцох брендүүд
  PUT    /pallets/sales-settings            — тохиргоо хадгалах (засах эрхтэй)
  GET    /pallets/sales-export?tag=&status=&q= — борлуулалтын Excel (хайрцаг, дүн, поддон)

Дүнгийн томьёо (calc):
  boxes_per_pallet = boxes_per_layer × layers
  pcs_per_box      = override > 0 ? override : Product.pack_ratio
  pcs_per_pallet   = boxes_per_pallet × pcs_per_box
  unit_weight_kg   = override > 0 ? override : Product.unit_weight   (ширхгийн хувийн жин)
  box_weight_kg    = override > 0 ? override : unit_weight_kg × pcs_per_box
  pallet_weight_kg = boxes_per_pallet × box_weight_kg   (+ поддоны жин орохгүй)
  total_height_cm  = template.height_cm + layers × box_height_cm   (шалнаас дээд хайрцаг)
  footprint        = template.length × width; нэг үеийн талбай = boxes_per_layer × box_L × box_W
  fit warnings     — үеийн талбай > поддоны талбай, өндөр/даац хэтэрсэн
Эрх: харах/хайх — нэвтэрсэн хэн ч; засах — admin/supervisor/manager/warehouse_clerk.
"""
from __future__ import annotations

import io
import re
import threading
import time
from datetime import datetime
from typing import NamedTuple, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import and_, or_, text
from sqlalchemy.orm import Session

from app.api.deps import get_current_user, get_db, require_role
from app.core.audit import audit
from app.models.pallet import PalletSalesSettings, PalletTemplate, ProductPallet
from app.models.product import Product
from app.models.user import User

router = APIRouter(prefix="/pallets", tags=["pallets"])
EDIT_ROLES = ("admin", "supervisor", "manager", "warehouse_clerk")
DEFAULT_TAG = "Архи Ус ундаа пиво"        # архины агуулах
SALES_MONTHS = (1, 2, 3, 6, 8)            # анхдагч сарын дундаж — тохиргооноос өөрчилнө (PalletSalesSettings)
DEFAULT_SINGLE_MONTH = 7
DEFAULT_PALLET_BRANDS = "Апу;Тотал ус ундаа;Тотал архи"


def _settings(db: Session) -> PalletSalesSettings:
    """Singleton тохиргоо (id=1) — байхгүй бол анхдагчаар үүсгэнэ."""
    st = db.get(PalletSalesSettings, 1)
    if st is None:
        st = PalletSalesSettings(id=1, avg_months=",".join(map(str, SALES_MONTHS)), single_month=DEFAULT_SINGLE_MONTH,
                                 pallet_brands=DEFAULT_PALLET_BRANDS, updated_by="", updated_at=datetime.utcnow())
        db.add(st)
        db.commit()
        db.refresh(st)
    return st


def _parse_months(raw: str) -> tuple[int, ...]:
    out = sorted({int(x) for x in re.findall(r"\d+", raw or "") if 1 <= int(x) <= 12})
    return tuple(out) or SALES_MONTHS


def _avg_months(db: Session) -> tuple[int, ...]:
    """Сарын дунджид орох сарууд — цонх, сард поддон, Excel бүгд үүнийг ашиглана."""
    return _parse_months(_settings(db).avg_months)


def _brand_list(raw: str) -> list[str]:
    return [b.strip() for b in re.split(r"[;\n]", raw or "") if b.strip()]


def _iso(d):
    return d.isoformat() if d else None


def _tpl_dict(t: PalletTemplate, used: int = 0) -> dict:
    return {"id": t.id, "name": t.name, "length_cm": t.length_cm, "width_cm": t.width_cm,
            "height_cm": t.height_cm, "max_weight_kg": t.max_weight_kg, "max_height_cm": t.max_height_cm,
            "note": t.note or "", "created_at": _iso(t.created_at), "used_by": used}


def _product_dict(p: Product) -> dict:
    return {"item_code": p.item_code, "name": p.name or "", "brand": p.brand or "",
            "warehouse_name": p.warehouse_name or "", "barcode": p.barcode or "",
            "pack_ratio": float(p.pack_ratio or 0), "unit_weight": float(p.unit_weight or 0),
            "box_weight_kg": round(float(p.unit_weight or 0) * float(p.pack_ratio or 0), 3)}


def calc(cfg: ProductPallet, tpl: Optional[PalletTemplate], p: Optional[Product]) -> dict:
    pcs_per_box = cfg.pcs_per_box_override if cfg.pcs_per_box_override > 0 else float((p.pack_ratio if p else 0) or 0)
    unit_w = cfg.unit_weight_kg_override if cfg.unit_weight_kg_override > 0 else float((p.unit_weight if p else 0) or 0)
    box_w = cfg.box_weight_kg_override if cfg.box_weight_kg_override > 0 else unit_w * pcs_per_box
    boxes = int(cfg.boxes_per_layer or 0) * int(cfg.layers or 0)
    th = (tpl.height_cm if tpl else 0.0) + (cfg.layers or 0) * (cfg.box_height_cm or 0)
    layer_area = (cfg.boxes_per_layer or 0) * (cfg.box_length_cm or 0) * (cfg.box_width_cm or 0)
    pallet_area = (tpl.length_cm * tpl.width_cm) if tpl else 0.0
    warnings = []
    if tpl and pallet_area > 0 and layer_area > pallet_area * 1.0001:
        warnings.append(f"Нэг үеийн хайрцгийн талбай ({layer_area:.0f} см²) поддоны талбайгаас ({pallet_area:.0f} см²) их байна")
    if tpl and tpl.max_height_cm > 0 and th > tpl.max_height_cm:
        warnings.append(f"Нийт өндөр {th:.0f} см > зөвшөөрөгдөх {tpl.max_height_cm:.0f} см")
    weight = boxes * box_w
    if tpl and tpl.max_weight_kg > 0 and weight > tpl.max_weight_kg:
        warnings.append(f"Жин {weight:.0f} кг > даац {tpl.max_weight_kg:.0f} кг")
    if pcs_per_box <= 0:
        warnings.append("Хайрцаг дахь ширхэг мэдэгдэхгүй (мастерт 0) — гараар оруулна уу")
    if box_w <= 0:
        warnings.append("Хайрцагны жин мэдэгдэхгүй (мастерт жин 0) — гараар оруулна уу")
    return {
        "boxes_per_pallet": boxes, "pcs_per_box": pcs_per_box, "pcs_per_pallet": round(boxes * pcs_per_box, 3),
        "unit_weight_kg": round(unit_w, 4), "box_weight_kg": round(box_w, 3), "pallet_weight_kg": round(weight, 2),
        "total_height_cm": round(th, 1), "stack_height_cm": round((cfg.layers or 0) * (cfg.box_height_cm or 0), 1),
        "pallet_length_cm": tpl.length_cm if tpl else 0, "pallet_width_cm": tpl.width_cm if tpl else 0,
        "layer_area_cm2": round(layer_area, 1), "pallet_area_cm2": round(pallet_area, 1),
        "area_fill_pct": round(layer_area / pallet_area * 100, 1) if pallet_area > 0 else None,
        "warnings": warnings,
    }


def _cfg_dict(cfg: ProductPallet, tpl: Optional[PalletTemplate], p: Optional[Product]) -> dict:
    return {
        "item_code": cfg.item_code, "template_id": cfg.template_id, "template_name": tpl.name if tpl else "",
        "box_length_cm": cfg.box_length_cm, "box_width_cm": cfg.box_width_cm, "box_height_cm": cfg.box_height_cm,
        "boxes_per_layer": cfg.boxes_per_layer, "layers": cfg.layers,
        "pcs_per_box_override": cfg.pcs_per_box_override, "unit_weight_kg_override": cfg.unit_weight_kg_override,
        "box_weight_kg_override": cfg.box_weight_kg_override,
        "note": cfg.note or "", "updated_by": cfg.updated_by or "", "updated_at": _iso(cfg.updated_at),
        "product": _product_dict(p) if p else {"item_code": cfg.item_code, "name": "", "pack_ratio": 0, "unit_weight": 0},
        "calc": calc(cfg, tpl, p),
    }


# ── Загвар ───────────────────────────────────────────────────────────────────

class TemplateIn(BaseModel):
    name: str
    length_cm: float = 120
    width_cm: float = 80
    height_cm: float = 15
    max_weight_kg: float = 0
    max_height_cm: float = 0
    note: str = ""


@router.get("/templates")
def list_templates(db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    used = {}
    for tid, in db.query(ProductPallet.template_id).all():
        used[tid] = used.get(tid, 0) + 1
    return [_tpl_dict(t, used.get(t.id, 0)) for t in db.query(PalletTemplate).order_by(PalletTemplate.name).all()]


@router.post("/templates")
def create_template(body: TemplateIn, request: Request, db: Session = Depends(get_db),
                    u: User = Depends(require_role(*EDIT_ROLES))):
    if not body.name.strip():
        raise HTTPException(400, "Нэр оруулна уу.")
    if body.length_cm <= 0 or body.width_cm <= 0:
        raise HTTPException(400, "Урт, өргөн 0-ээс их байх ёстой.")
    t = PalletTemplate(name=body.name.strip(), length_cm=body.length_cm, width_cm=body.width_cm,
                       height_cm=max(body.height_cm, 0), max_weight_kg=max(body.max_weight_kg, 0),
                       max_height_cm=max(body.max_height_cm, 0), note=body.note.strip()[:300])
    db.add(t); db.commit()
    audit(db, request, u, action="pallet_template_create", entity_type="pallet_template", entity_id=t.id,
          extra=body.model_dump(), autocommit=True)
    return _tpl_dict(t)


@router.put("/templates/{tid}")
def update_template(tid: int, body: TemplateIn, request: Request, db: Session = Depends(get_db),
                    u: User = Depends(require_role(*EDIT_ROLES))):
    t = db.query(PalletTemplate).filter(PalletTemplate.id == tid).first()
    if not t:
        raise HTTPException(404, "Загвар олдсонгүй.")
    if not body.name.strip() or body.length_cm <= 0 or body.width_cm <= 0:
        raise HTTPException(400, "Нэр, урт, өргөн зөв оруулна уу.")
    t.name, t.length_cm, t.width_cm = body.name.strip(), body.length_cm, body.width_cm
    t.height_cm, t.max_weight_kg, t.max_height_cm = max(body.height_cm, 0), max(body.max_weight_kg, 0), max(body.max_height_cm, 0)
    t.note = body.note.strip()[:300]
    db.commit()
    audit(db, request, u, action="pallet_template_update", entity_type="pallet_template", entity_id=t.id,
          extra=body.model_dump(), autocommit=True)
    return _tpl_dict(t)


@router.delete("/templates/{tid}")
def delete_template(tid: int, request: Request, db: Session = Depends(get_db),
                    u: User = Depends(require_role(*EDIT_ROLES))):
    t = db.query(PalletTemplate).filter(PalletTemplate.id == tid).first()
    if not t:
        return {"ok": True}
    n = db.query(ProductPallet).filter(ProductPallet.template_id == tid).count()
    if n:
        raise HTTPException(400, f"{n} бараа энэ загварыг ашиглаж байна — эхлээд тэдгээрийг өөр загварт шилжүүлнэ үү.")
    db.delete(t); db.commit()
    audit(db, request, u, action="pallet_template_delete", entity_type="pallet_template", entity_id=tid, autocommit=True)
    return {"ok": True}


# ── Бараа ────────────────────────────────────────────────────────────────────

@router.get("/lookup")
def lookup(q: str = Query(..., min_length=1, max_length=80), tag: str = Query(DEFAULT_TAG, max_length=120),
           db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    """Баркод / код → бараа (мастер) + одоогийн поддоны тохиргоо (байвал) + сарын борлуулалт."""
    from app.services.hall_count import resolve_product
    p = resolve_product(db, q)
    if not p or not p.item_code:
        return {"query": q, "found": False, "product": None, "config": None}
    cfg = db.query(ProductPallet).filter(ProductPallet.item_code == p.item_code).first()
    tpl = db.query(PalletTemplate).filter(PalletTemplate.id == cfg.template_id).first() if cfg else None
    try:
        sales = _product_sales(db, p.item_code, tag.strip())
    except Exception as e:                                        # борлуулалтгүй ч тохиргоо оруулж болно
        print(f"[pallets] sales алдаа: {e}")
        sales = None
    return {"query": q, "found": True, "product": _product_dict(p),
            "config": _cfg_dict(cfg, tpl, p) if cfg else None, "sales": sales}


@router.get("/products")
def list_products(q: str = Query("", max_length=80), template_id: Optional[int] = None,
                  db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    qry = db.query(ProductPallet)
    if template_id:
        qry = qry.filter(ProductPallet.template_id == template_id)
    cfgs = qry.order_by(ProductPallet.updated_at.desc()).all()
    codes = [c.item_code for c in cfgs]
    prods = {}
    for i in range(0, len(codes), 900):
        for p in db.query(Product).filter(Product.item_code.in_(codes[i:i + 900])).all():
            prods[p.item_code] = p
    tpls = {t.id: t for t in db.query(PalletTemplate).all()}
    out = []
    ql = q.strip().lower()
    for c in cfgs:
        p = prods.get(c.item_code)
        if ql and ql not in c.item_code.lower() and ql not in (p.name if p else "").lower():
            continue
        out.append(_cfg_dict(c, tpls.get(c.template_id), p))
    return {"total": len(out), "items": out}


class ProductPalletIn(BaseModel):
    template_id: int
    box_length_cm: float = 0
    box_width_cm: float = 0
    box_height_cm: float = 0
    boxes_per_layer: int = 0
    layers: int = 0
    pcs_per_box_override: float = 0
    unit_weight_kg_override: float = 0
    box_weight_kg_override: float = 0
    note: str = ""


@router.put("/products/{item_code}")
def upsert_product(item_code: str, body: ProductPalletIn, request: Request, db: Session = Depends(get_db),
                   u: User = Depends(require_role(*EDIT_ROLES))):
    code = re.sub(r"\s+", "", item_code.strip())
    p = db.query(Product).filter(Product.item_code == code).first()
    if not p:
        raise HTTPException(404, f"«{code}» бараа мастерт олдсонгүй.")
    tpl = db.query(PalletTemplate).filter(PalletTemplate.id == body.template_id).first()
    if not tpl:
        raise HTTPException(400, "Поддоны загвар сонгоно уу.")
    if body.boxes_per_layer <= 0 or body.layers <= 0:
        raise HTTPException(400, "Нэг үеийн хайрцаг ба үеийн тоо 0-ээс их байх ёстой.")
    if body.box_height_cm <= 0:
        raise HTTPException(400, "Хайрцагны өндөр оруулна уу (нийт өндөр бодоход хэрэгтэй).")
    cfg = db.query(ProductPallet).filter(ProductPallet.item_code == code).first()
    if not cfg:
        cfg = ProductPallet(item_code=code)
        db.add(cfg)
    cfg.template_id = tpl.id
    cfg.box_length_cm, cfg.box_width_cm, cfg.box_height_cm = max(body.box_length_cm, 0), max(body.box_width_cm, 0), body.box_height_cm
    cfg.boxes_per_layer, cfg.layers = int(body.boxes_per_layer), int(body.layers)
    cfg.pcs_per_box_override = max(body.pcs_per_box_override, 0)
    cfg.unit_weight_kg_override = max(body.unit_weight_kg_override, 0)
    cfg.box_weight_kg_override = max(body.box_weight_kg_override, 0)
    cfg.note = body.note.strip()[:300]
    cfg.updated_by = str(getattr(u, "nickname", "") or getattr(u, "username", "") or "")
    cfg.updated_at = datetime.utcnow()
    db.commit()
    audit(db, request, u, action="product_pallet_save", entity_type="product_pallet", entity_id=cfg.id,
          extra={"item_code": code, **body.model_dump()}, autocommit=True)
    return _cfg_dict(cfg, tpl, p)


@router.delete("/products/{item_code}")
def delete_product(item_code: str, request: Request, db: Session = Depends(get_db),
                   u: User = Depends(require_role(*EDIT_ROLES))):
    cfg = db.query(ProductPallet).filter(ProductPallet.item_code == item_code.strip()).first()
    if cfg:
        db.delete(cfg); db.commit()
        audit(db, request, u, action="product_pallet_delete", entity_type="product_pallet",
              extra={"item_code": item_code}, autocommit=True)
    return {"ok": True}


# ── Хуулах: ижил хэмжээтэй өөр кодтой бараанд ──────────────────────────────────

def _same(a: float, b: float) -> bool:
    a, b = float(a or 0), float(b or 0)
    return abs(a - b) <= max(abs(a), abs(b)) * 0.01 + 1e-9


_SIZE_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*(мл|ml|литр|л|l|гр|gr|г|g|кг|kg)(?![a-zа-яөү])", re.I)
_SIZE_MUL = {"мл": 1, "ml": 1, "литр": 1000, "л": 1000, "l": 1000, "гр": 1, "gr": 1, "г": 1, "g": 1, "кг": 1000, "kg": 1000}


def _pack_size(name: str) -> tuple[str, float, str] | None:
    """Нэрнээс савлагааны хэмжээ: «Coca-Cola 1.5л» → ('ml', 1500, '1.5л'). Сүүлчийн таарцыг авна."""
    ms = list(_SIZE_RE.finditer(name or ""))
    if not ms:
        return None
    m = ms[-1]
    unit = m.group(2).lower()
    kind = "ml" if unit in ("мл", "ml", "литр", "л", "l") else "g"
    return kind, round(float(m.group(1).replace(",", ".")) * _SIZE_MUL[unit], 3), m.group(0).replace(" ", "")


def _layout(cfg: ProductPallet) -> tuple:
    return (cfg.template_id, cfg.box_length_cm, cfg.box_width_cm, cfg.box_height_cm, cfg.boxes_per_layer, cfg.layers)


# Хуулах талбарууд: түлхүүр → ProductPallet баганууд
COPY_FIELDS = {
    "template": ("template_id",),
    "box_dims": ("box_length_cm", "box_width_cm", "box_height_cm"),
    "boxes_per_layer": ("boxes_per_layer",),
    "layers": ("layers",),
    "pcs_per_box": ("pcs_per_box_override",),
    "unit_weight": ("unit_weight_kg_override",),
    "box_weight": ("box_weight_kg_override",),
    "note": ("note",),
}
# Тохиргоогүй бараанд шинээр үүсгэхэд заавал хэрэгтэй (загвар, өндөр, үе байхгүй бол тохиргоо хүчингүй)
REQUIRED_NEW = ("template", "box_dims", "boxes_per_layer", "layers")
OVERRIDE_FIELDS = ("pcs_per_box", "unit_weight", "box_weight")


class _PRow(NamedTuple):
    item_code: str
    name: str
    name_l: str
    brand: str
    brand_code: str
    barcode: str
    barcode_l: str
    pack_ratio: float
    unit_weight: float
    warehouse_name: str
    size: tuple | None
    hay: str                     # код + нэр + баркод (жижиг үсгээр) — хайлтад


# Мастерын ~22мянган барааг хүсэлт бүрд ORM-оор уншвал ~0.5с — нэг удаа индексжүүлж кэшилнэ.
# Мастер шинэчлэгдэхэд (тоо, max id, ш/хайрцаг·жингийн нийлбэр өөрчлөгдөхөд) эсвэл 5 минут тутам дахин уншина.
_PIDX: dict = {"sig": None, "at": 0.0, "rows": []}
_PIDX_LOCK = threading.Lock()


def _product_index(db: Session) -> list[_PRow]:
    sig = tuple(db.execute(text("SELECT count(*), max(id), total(pack_ratio), total(unit_weight) FROM products")).one())
    if _PIDX["sig"] == sig and time.monotonic() - _PIDX["at"] < 300:
        return _PIDX["rows"]
    with _PIDX_LOCK:
        if _PIDX["sig"] == sig and time.monotonic() - _PIDX["at"] < 300:
            return _PIDX["rows"]
        cur = db.connection().connection.cursor()
        cur.execute("SELECT item_code, name, brand, brand_code, barcode, pack_ratio, unit_weight, warehouse_name FROM products")
        rows, seen = [], set()
        for code, name, brand, bcode, barcode, pr, uw, wh in cur.fetchall():
            if not code or code in seen:
                continue
            seen.add(code)
            name = name or ""
            rows.append(_PRow(code, name, name.lower(), brand or "", bcode or "", barcode or "", (barcode or "").lower(),
                              float(pr or 0), float(uw or 0), wh or "", _pack_size(name),
                              f"{code} {name} {barcode or ''}".lower()))
        rows.sort(key=lambda r: r.name_l)
        _PIDX.update(sig=sig, at=time.monotonic(), rows=rows)
        return rows


def _row_dict(r: _PRow) -> dict:
    return {"item_code": r.item_code, "name": r.name, "brand": r.brand, "warehouse_name": r.warehouse_name,
            "barcode": r.barcode, "pack_ratio": r.pack_ratio, "unit_weight": r.unit_weight,
            "box_weight_kg": round(r.unit_weight * r.pack_ratio, 3)}


@router.get("/copy-candidates")
def copy_candidates(source: str = Query(..., min_length=1, max_length=64), q: str = Query("", max_length=80),
                    scope: str = Query("all", pattern="^(all|brand|match)$"), unconfigured: bool = False,
                    offset: int = Query(0, ge=0), limit: int = Query(100, ge=1, le=300),
                    db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    """Хуулах бараа сонгох жагсаалт — мастерын БҮХ бараа (scope=all), эх барааны брэнд (brand),
    эсвэл нэрэн дэх савлагааны хэмжээ + ш/хайрцаг ижил бараа (match). q: код/нэр/баркод."""
    from difflib import SequenceMatcher
    from app.services.hall_count import resolve_product
    src = db.query(Product).filter(Product.item_code == source.strip()).first()
    if not src:
        raise HTTPException(404, "Эх бараа олдсонгүй.")
    src_cfg = db.query(ProductPallet).filter(ProductPallet.item_code == src.item_code).first()
    qs = q.strip()
    ql = qs.lower()
    exact = ""
    if qs and re.fullmatch(r"[0-9A-Za-z_\-]{3,}", qs):          # код/баркод шиг бол яг таарцыг олно
        hit = resolve_product(db, qs)
        if hit and hit.item_code and hit.id:
            exact = hit.item_code
    cfgs = {c.item_code: c for c in db.query(ProductPallet).all()}
    ssize = _pack_size(src.name)
    toks = ql.split()
    found = []
    for r in _product_index(db):
        code = r.item_code
        if code == src.item_code:
            continue
        if scope == "brand" and not ((r.brand_code == src.brand_code) if src.brand_code else (bool(src.brand) and r.brand == src.brand)):
            continue
        if toks and code != exact and not all(t in r.hay for t in toks):   # үг бүр (AND): «түмний 2л»
            continue
        c = cfgs.get(code)
        if unconfigured and c:
            continue
        same_pack, same_weight = _same(r.pack_ratio, src.pack_ratio), _same(r.unit_weight, src.unit_weight)
        psize = r.size
        same_size = bool(ssize and psize and ssize[:2] == psize[:2])
        # Ижил хэмжээтэй: нэрэн дэх савлагааны хэмжээ + ш/хайрцаг (нэрэнд хэмжээгүй бол ш/хайрцаг + жин)
        match = (same_size and same_pack) if ssize else (same_pack and same_weight)
        if scope == "match" and not match:
            continue
        found.append((r, c, same_pack, same_weight, psize, same_size, match))
    if scope == "all":
        # индекс нэрээр эрэмбэлэгдсэн — яг таарсан, кодоор эхэлсэн, үгийн эхэнд таарсан нь эхэнд (sort тогтвортой)
        wre = re.compile(r"(?:^|[\s/\-(\"«])" + re.escape(toks[0])) if toks else None
        found.sort(key=lambda x: (x[0].item_code != exact, bool(ql) and not x[0].item_code.lower().startswith(ql),
                                  bool(wre) and not wre.search(x[0].name_l)))
    else:
        sname = (src.name or "").lower()
        def score(x):
            r, c, sp, sw, _, ss, _m = x
            sim = SequenceMatcher(None, sname, r.name_l).ratio()
            return (r.item_code != exact, -(sim + 0.6 * ss + 0.25 * sp + 0.15 * sw - 0.5 * (c is not None)))
        found.sort(key=score)
    tpls = {t.id: t for t in db.query(PalletTemplate).all()}
    out = []
    for r, c, same_pack, same_weight, psize, same_size, match in found[offset:offset + limit]:
        out.append({
            **_row_dict(r), "exact": r.item_code == exact, "same_pack": same_pack, "same_weight": same_weight,
            "size": psize[2] if psize else "", "same_size": same_size, "match": match,
            "config": ({"template_id": c.template_id, "template_name": tpls[c.template_id].name if c.template_id in tpls else "",
                        "box_length_cm": c.box_length_cm, "box_width_cm": c.box_width_cm, "box_height_cm": c.box_height_cm,
                        "boxes_per_layer": c.boxes_per_layer, "layers": c.layers,
                        "pcs_per_box_override": c.pcs_per_box_override, "unit_weight_kg_override": c.unit_weight_kg_override,
                        "box_weight_kg_override": c.box_weight_kg_override, "note": c.note or "",
                        "same_as_source": bool(src_cfg) and _layout(c) == _layout(src_cfg)} if c else None),
        })
    return {"source": {**_product_dict(src), "size": ssize[2] if ssize else ""}, "scope": scope,
            "total": len(found), "offset": offset, "items": out}


class CopyIn(BaseModel):
    targets: list[str]
    fields: Optional[list[str]] = None  # COPY_FIELDS-ийн түлхүүрүүд; None бол загвар+хэмжээ+өрөлт
    include_overrides: bool = False     # (хуучин) fields өгөөгүй үед ш/хайрцаг, жингийн засварыг нэмнэ
    overwrite: bool = False             # тохиргоотой барааны сонгосон утгуудыг солих


@router.post("/products/{item_code}/copy")
def copy_product(item_code: str, body: CopyIn, request: Request, db: Session = Depends(get_db),
                 u: User = Depends(require_role(*EDIT_ROLES))):
    """Эх барааны поддоны тохиргооноос СОНГОСОН талбаруудыг бусад бараанд хуулна.
    Тохиргоогүй бараанд шинээр үүсгэхэд загвар, хайрцагны хэмжээ, нэг үеийн хайрцаг, үе заавал;
    тохиргоотой бараанд overwrite үед зөвхөн сонгосон талбарууд солигдоно (бусад нь хэвээр)."""
    src_code = re.sub(r"\s+", "", item_code.strip())
    src = db.query(ProductPallet).filter(ProductPallet.item_code == src_code).first()
    if not src:
        raise HTTPException(404, "Эх барааны поддоны тохиргоо хадгалагдаагүй байна — эхлээд хадгална уу.")
    if body.fields is None:
        fields = list(REQUIRED_NEW) + (list(OVERRIDE_FIELDS) if body.include_overrides else [])
    else:
        fields = [f for f in COPY_FIELDS if f in set(body.fields)]
    if not fields:
        raise HTTPException(400, "Юуг хуулахаа сонгоно уу.")
    full_new = all(f in fields for f in REQUIRED_NEW)
    targets = []
    for t in body.targets:
        c = re.sub(r"\s+", "", str(t or ""))
        if c and c != src_code and c not in targets:
            targets.append(c)
    if not targets:
        raise HTTPException(400, "Хуулах бараа сонгоно уу.")
    if len(targets) > 300:
        raise HTTPException(400, "Нэг удаад 300 хүртэл бараа хуулна.")
    prods = {p.item_code: p for p in db.query(Product).filter(Product.item_code.in_(targets)).all()}
    have = {c.item_code: c for c in db.query(ProductPallet).filter(ProductPallet.item_code.in_(targets)).all()}
    tpls = {t.id: t for t in db.query(PalletTemplate).all()}
    who = str(getattr(u, "nickname", "") or getattr(u, "username", "") or "")
    now = datetime.utcnow()
    copied, skipped = [], []
    for code in targets:
        p = prods.get(code)
        if not p:
            skipped.append({"item_code": code, "reason": "Мастерт олдсонгүй"}); continue
        cfg = have.get(code)
        if cfg and not body.overwrite:
            skipped.append({"item_code": code, "name": p.name or "", "reason": "Тохиргоотой (солихыг сонгоогүй)"}); continue
        if not cfg:
            if not full_new:
                skipped.append({"item_code": code, "name": p.name or "",
                                "reason": "Поддоны тохиргоогүй — загвар, хайрцагны хэмжээ, нэг үеийн хайрцаг, үеийг мөн хуулах шаардлагатай"})
                continue
            cfg = ProductPallet(item_code=code, pcs_per_box_override=0.0, unit_weight_kg_override=0.0,
                                box_weight_kg_override=0.0, note="")
            db.add(cfg)
        for f in fields:
            for attr in COPY_FIELDS[f]:
                setattr(cfg, attr, getattr(src, attr))
        if "note" not in fields and not (cfg.note or "").strip():
            cfg.note = f"{src_code}-аас хуулсан"
        cfg.updated_by, cfg.updated_at = who, now
        copied.append((cfg, p))
    db.commit()
    if copied:
        audit(db, request, u, action="product_pallet_copy", entity_type="product_pallet", entity_id=src.id,
              extra={"source": src_code, "fields": fields, "targets": [c.item_code for c, _ in copied],
                     "skipped": skipped, "overwrite": body.overwrite}, autocommit=True)
    return {"source": src_code, "fields": fields,
            "copied": [_cfg_dict(c, tpls.get(c.template_id), p) for c, p in copied], "skipped": skipped}


# ── Ажлын жагсаалт: байршлын tag-ийн бүх бараа борлуулалтын эрэмбээр ──────────────

_SALES: dict = {"key": None, "slots": [], "map": {}}
_SALES_LOCK = threading.Lock()


def _sales(db: Session, months: tuple = SALES_MONTHS) -> tuple[list[tuple[int, int]], dict]:
    """(slots, {code: {(year, month): qty}}) — сар бүрийн ХАМГИЙН СҮҮЛИЙН (дататай) жилийг авна.
    qty = агуулах + заалны борлуулалт (ширхэг). Сарын борлуулалт шинэчлэгдэхэд кэш шинэчлэгдэнэ."""
    from app.models.product_monthly_sales import ProductMonthlySales as PMS
    sig = tuple(db.execute(text("SELECT count(*), max(updated_at) FROM product_monthly_sales")).one())
    key = (sig, tuple(months))
    if _SALES["key"] == key:
        return _SALES["slots"], _SALES["map"]
    with _SALES_LOCK:
        if _SALES["key"] == key:
            return _SALES["slots"], _SALES["map"]
        avail = db.query(PMS.year, PMS.month).distinct().all()
        slots = sorted((max(y for y, mm in avail if mm == m), m) for m in months if any(mm == m for _, mm in avail))
        out: dict = {}
        if slots:
            cond = or_(*[and_(PMS.year == y, PMS.month == m) for y, m in slots])
            for code, y, m, w, s in db.query(PMS.item_code, PMS.year, PMS.month, PMS.qty_warehouse, PMS.qty_showroom).filter(cond):
                out.setdefault(code, {})[(y, m)] = float(w or 0) + float(s or 0)
        _SALES.update(key=key, slots=slots, map=out)
        return slots, out


def _master_tag_map() -> dict[str, list[str]]:
    from app.api.tag_location_check import _get_master_tags
    return _get_master_tags()


def _worklist_rows(db: Session, tag: str, status: str = "all", q: str = "", template_id: Optional[int] = None):
    """Tag-ийн бүх бараа (мастер), сарын дундаж борлуулалтаар буурах эрэмбэтэй.
    → (rows, counts, slots). rank нь шүүлтээс өмнөх (tag доторх) эрэмбэ."""
    mt = _master_tag_map()
    codes = [c for c, tags in mt.items() if not tag or tag in tags]
    idx = {r.item_code: r for r in _product_index(db)}
    slots, sales = _sales(db, _avg_months(db))
    cfgs = {c.item_code: c for c in db.query(ProductPallet).all()}
    tpls = {t.id: t for t in db.query(PalletTemplate).all()}
    base = []
    for code in codes:
        per = sales.get(code, {})
        avg = sum(per.get(s, 0.0) for s in slots) / len(slots) if slots else 0.0
        base.append((code, avg, per))
    base.sort(key=lambda x: (-x[1], x[0]))
    done = sum(1 for c, _, _ in base if c in cfgs)
    counts = {"total": len(base), "done": done, "todo": len(base) - done}
    toks = q.strip().lower().split()
    rows = []
    for rank, (code, avg, per) in enumerate(base, 1):
        c = cfgs.get(code)
        if (status == "done" and not c) or (status == "todo" and c):
            continue
        if template_id and (not c or c.template_id != template_id):
            continue
        r = idx.get(code)
        if toks and not all(t in (r.hay if r else code.lower()) for t in toks):
            continue
        cfg = None
        if c:
            tpl = tpls.get(c.template_id)
            k = calc(c, tpl, r)
            ppp = k["pcs_per_pallet"]
            cfg = {"template_id": c.template_id, "template_name": tpl.name if tpl else "",
                   "box_length_cm": c.box_length_cm, "box_width_cm": c.box_width_cm, "box_height_cm": c.box_height_cm,
                   "boxes_per_layer": c.boxes_per_layer, "layers": c.layers, "note": c.note or "",
                   "updated_by": c.updated_by or "", "updated_at": _iso(c.updated_at), "calc": k,
                   "pallets_per_month": round(avg / ppp, 2) if ppp > 0 else None}
        rows.append({
            "rank": rank, "item_code": code, "name": r.name if r else "", "brand": r.brand if r else "",
            "barcode": r.barcode if r else "", "pack_ratio": r.pack_ratio if r else 0.0,
            "unit_weight": r.unit_weight if r else 0.0, "tags": mt.get(code, []),
            "avg_monthly": round(avg, 1), "months": [round(per.get(s, 0.0), 1) for s in slots],
            "configured": c is not None, "config": cfg,
        })
    return rows, counts, slots


def warm_worklist_caches() -> None:
    """main-ийн warm loop-оос (60с тутам) — анхны нээлт мастер/борлуулалт parse хүлээхгүй."""
    from app.core.db import SessionLocal
    db = SessionLocal()
    try:
        _product_index(db)
        _sales(db, _avg_months(db))
        _master_tag_map()
    finally:
        db.close()


def _slot_list(slots) -> list[dict]:
    return [{"year": y, "month": m} for y, m in slots]


@router.get("/worklist")
def worklist(tag: str = Query(DEFAULT_TAG, max_length=120), status: str = Query("all", pattern="^(all|todo|done)$"),
             q: str = Query("", max_length=80), offset: int = Query(0, ge=0), limit: int = Query(200, ge=1, le=2000),
             db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    """Поддон мэдээлэл оруулах ажлын жагсаалт — tag-ийн бүх бараа борлуулалтын эрэмбээр,
    оруулсан/оруулаагүйгээр. tag="" бол мастерын бүх бараа."""
    rows, counts, slots = _worklist_rows(db, tag.strip(), status, q)
    from collections import Counter
    tc = Counter(t for tags in _master_tag_map().values() for t in tags)
    return {"tag": tag.strip(), "tags": [{"name": k, "count": v} for k, v in tc.most_common()],
            "months": _slot_list(slots), "counts": counts, "total": len(rows), "offset": offset,
            "items": rows[offset:offset + limit]}


def _product_sales(db: Session, code: str, tag: str) -> dict:
    """Нэг барааны сарын борлуулалт + tag доторх борлуулалтын эрэмбэ (поддон оруулах цонхонд)."""
    slots, sales = _sales(db, _avg_months(db))
    per = sales.get(code, {})
    avg = sum(per.get(s, 0.0) for s in slots) / len(slots) if slots else 0.0
    tags = _master_tag_map().get(code, [])
    use = tag if tag in tags else (DEFAULT_TAG if DEFAULT_TAG in tags else (tags[0] if tags else ""))
    rank = total = None
    if use:
        rows, counts, _ = _worklist_rows(db, use)
        total = counts["total"]
        rank = next((r["rank"] for r in rows if r["item_code"] == code), None)
    return {"months": [{"year": y, "month": m, "qty": round(per.get((y, m), 0.0), 1)} for y, m in slots],
            "avg_monthly": round(avg, 1), "tag": use, "tags": tags, "rank": rank, "rank_total": total}


# ── Excel ────────────────────────────────────────────────────────────────────

@router.get("/export")
def export_xlsx(template_id: Optional[int] = None, tag: Optional[str] = Query(None, max_length=120),
                status: str = Query("all", pattern="^(all|todo|done)$"),
                db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    """tag өгвөл тухайн tag-ийн бүх бараа (оруулсан/оруулаагүй) борлуулалтын эрэмбээр;
    өгөхгүй бол мастерын тохиргоотой бүх бараа."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if tag is None:
        data, _, slots = _worklist_rows(db, "", "done", template_id=template_id)
    else:
        data, _, slots = _worklist_rows(db, tag.strip(), status, template_id=template_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Поддон хураалт"
    mcols = [f"{y}-{m:02d} борлуулалт" for y, m in slots]
    cols = (["Эрэмбэ", "Код", "Нэр", "Бренд", "Байршил tag", "Поддон мэдээлэл", "Сарын дундаж борлуулалт (ш)"] + mcols
            + ["Сард поддон", "Поддон загвар", "Поддон урт (см)", "Поддон өргөн (см)", "Поддон өндөр (см)",
               "Хайрцаг урт (см)", "Хайрцаг өргөн (см)", "Хайрцаг өндөр (см)", "Нэг үед (хайрцаг)", "Үе", "Нийт хайрцаг",
               "Ширхэг/хайрцаг", "Нийт ширхэг", "Хувийн жин (кг/ш)", "Хайрцагны жин (кг)", "Поддоны бараа жин (кг)",
               "Өрөлтийн өндөр (см)", "Шалнаас дээд хайрцаг (см)", "Талбайн дүүргэлт %", "Анхааруулга", "Тэмдэглэл",
               "Шинэчилсэн", "Хэн"])
    widths = ([7, 10, 40, 18, 18, 12, 12] + [11] * len(mcols)
              + [9, 18, 10, 10, 10, 10, 10, 10, 10, 6, 10, 10, 10, 12, 12, 14, 12, 14, 10, 40, 30, 16, 12])
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hf, hfont = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True)
    ok_fill, todo_fill = PatternFill("solid", fgColor="E2EFDA"), PatternFill("solid", fgColor="FFF2CC")
    for ci, h in enumerate(cols, 1):
        c = ws.cell(1, ci, h); c.fill = hf; c.font = hfont; c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 42
    status_col = 6
    for ri, d in enumerate(data, 2):
        g = d["config"]
        head = [d["rank"], d["item_code"], d["name"], d["brand"], ", ".join(d["tags"]),
                "Оруулсан" if g else "Оруулаагүй", d["avg_monthly"]] + d["months"]
        if g:
            k = g["calc"]
            rest = [g["pallets_per_month"], g["template_name"], k["pallet_length_cm"], k["pallet_width_cm"],
                    k["total_height_cm"] - k["stack_height_cm"], g["box_length_cm"], g["box_width_cm"], g["box_height_cm"],
                    g["boxes_per_layer"], g["layers"], k["boxes_per_pallet"], k["pcs_per_box"], k["pcs_per_pallet"],
                    k["unit_weight_kg"], k["box_weight_kg"], k["pallet_weight_kg"], k["stack_height_cm"], k["total_height_cm"],
                    k["area_fill_pct"], "; ".join(k["warnings"]), g["note"], (g["updated_at"] or "")[:16].replace("T", " "),
                    g["updated_by"]]
        else:
            rest = [None] * 23
        for ci, v in enumerate(head + rest, 1):
            ws.cell(ri, ci, v).border = border
        ws.cell(ri, status_col).fill = ok_fill if g else todo_fill
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(data) + 1, 2)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.oddHeader.left.text = "&D &T"; ws.oddFooter.right.text = "Хуудас &P / &N"

    ws2 = wb.create_sheet("Загварууд")
    for ci, h in enumerate(["Нэр", "Урт (см)", "Өргөн (см)", "Өндөр (см)", "Даац (кг)", "Дээд өндөр (см)", "Тэмдэглэл", "Ашигласан бараа"], 1):
        c = ws2.cell(1, ci, h); c.fill = hf; c.font = hfont; c.border = border
    for ri, t in enumerate(list_templates(db=db, u=u), 2):
        for ci, v in enumerate([t["name"], t["length_cm"], t["width_cm"], t["height_cm"], t["max_weight_kg"] or "",
                                t["max_height_cm"] or "", t["note"], t["used_by"]], 1):
            ws2.cell(ri, ci, v).border = border
    for ci, w in enumerate([24, 10, 10, 10, 10, 14, 30, 14], 1):
        ws2.column_dimensions[get_column_letter(ci)].width = w

    buf = io.BytesIO(); wb.save(buf)
    fname = f"{datetime.now():%Y%m%d}_poddon_huraalt{'_' + re.sub(r'[^0-9A-Za-zА-Яа-яӨөҮүЁё]+', '_', tag.strip()) if tag else ''}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})


# ── Борлуулалтын тохиргоо ба Excel ───────────────────────────────────────────

class SalesSettingsIn(BaseModel):
    avg_months: list[int]
    single_month: int = DEFAULT_SINGLE_MONTH
    pallet_brands: list[str] = []


def _settings_dict(st: PalletSalesSettings) -> dict:
    return {"avg_months": list(_parse_months(st.avg_months)), "single_month": int(st.single_month or DEFAULT_SINGLE_MONTH),
            "pallet_brands": _brand_list(st.pallet_brands), "updated_by": st.updated_by or "",
            "updated_at": _iso(st.updated_at)}


@router.get("/sales-settings")
def get_sales_settings(db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    return _settings_dict(_settings(db))


@router.put("/sales-settings")
def put_sales_settings(body: SalesSettingsIn, request: Request, db: Session = Depends(get_db),
                       u: User = Depends(require_role(*EDIT_ROLES))):
    months = sorted({int(m) for m in body.avg_months if 1 <= int(m) <= 12})
    if not months:
        raise HTTPException(400, "Дор хаяж нэг сар сонгоно уу")
    if not 1 <= int(body.single_month) <= 12:
        raise HTTPException(400, "Тусгай сар 1-12 байна")
    st = _settings(db)
    before = _settings_dict(st)
    st.avg_months = ",".join(map(str, months))
    st.single_month = int(body.single_month)
    st.pallet_brands = ";".join(b.strip() for b in body.pallet_brands if b.strip())
    st.updated_by = getattr(u, "username", "") or ""
    st.updated_at = datetime.utcnow()
    db.commit()
    audit(db, request, u, action="pallet_sales_settings", entity_type="pallet_sales_settings", entity_id=1,
          extra={"before": before, "after": _settings_dict(st)}, autocommit=True)
    return _settings_dict(st)


def _round_unit(v: float) -> int:
    """Нэгжийн орноор тоймлоно (0.5 → дээш). Python-ы round() банкны дүрэмтэй (2.5 → 2) тул Decimal."""
    from decimal import ROUND_HALF_UP, Decimal
    return int(Decimal(str(v)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


@router.get("/sales-export")
def sales_export(tag: str = Query(DEFAULT_TAG, max_length=120), status: str = Query("all", pattern="^(all|todo|done)$"),
                 q: str = Query("", max_length=80), db: Session = Depends(get_db), u: User = Depends(get_current_user)):
    """Борлуулалтын Excel — цонхны шүүлтээр (tag, төлөв, хайлт).

    Багана: Код, Нэр, Бренд, Хайрцаг дахь ширхэг, Нэгж үнэ, сарын дундаж борлуулалт
    хайрцгаар (нэгжийн орноор тоймлосон) ба × нэгж үнэ, тусгай сарын (анхдагч 7-р сар)
    борлуулалт хайрцгаар ба × нэгж үнэ, поддон хэмжээ (хайрцаг/поддон). Тохиргооны
    брендүүдийн бараанд борлуулалт хэдэн поддон болохыг бичнэ; поддон мэдээлэлгүй бол
    «Поддон мэдээлэл байхгүй».
    Нэгж үнэ = барааны мастерын сүүлийн орлогын үнэ (last_purchase_price).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    st = _settings(db)
    avg_months = _parse_months(st.avg_months)
    single = int(st.single_month or DEFAULT_SINGLE_MONTH)
    brand_names = _brand_list(st.pallet_brands)
    brands = {b.lower() for b in brand_names}

    rows, _counts, avg_slots = _worklist_rows(db, tag.strip(), status, q)
    single_slots, single_sales = _sales(db, (single,))
    cfgs = {c.item_code: c for c in db.query(ProductPallet).all()}
    tpls = {t.id: t for t in db.query(PalletTemplate).all()}
    idx = {r.item_code: r for r in _product_index(db)}
    prices: dict[str, float] = {}
    for code, price in db.execute(text("SELECT item_code, last_purchase_price FROM products")).all():
        if code and float(price or 0) > 0 and code not in prices:
            prices[code] = float(price)

    avg_label = ", ".join(str(m) for _y, m in avg_slots) or ", ".join(map(str, avg_months))
    single_label = f"{single_slots[0][0]} оны {single}-р сар" if single_slots else f"{single}-р сар"
    no_pallet = "Поддон мэдээлэл байхгүй"
    cols = ["Код", "Нэр", "Бренд", "Хайрцаг дахь ширхэг", "Нэгж үнэ",
            f"Сарын дундаж ({avg_label}-р сар), хайрцаг", f"Сарын дундаж ({avg_label}-р сар) × нэгж үнэ",
            f"{single_label}, хайрцаг", f"{single_label} × нэгж үнэ",
            "Поддон хэмжээ (хайрцаг/поддон)", "Сарын дундаж, поддон", f"{single_label}, поддон"]
    widths = [11, 42, 18, 11, 12, 14, 16, 13, 16, 15, 14, 14]

    wb = Workbook()
    ws = wb.active
    ws.title = "Борлуулалт"
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hf, hfont = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True)
    pallet_fill, missing_fill = PatternFill("solid", fgColor="E2EFDA"), PatternFill("solid", fgColor="FCE4D6")
    for ci, h in enumerate(cols, 1):
        c = ws.cell(1, ci, h); c.fill = hf; c.font = hfont; c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 48

    tot_avg_amt = tot_single_amt = 0.0
    tot_avg_box = tot_single_box = 0
    ri = 1
    for d in rows:
        code = d["item_code"]
        cfg = cfgs.get(code)
        r = idx.get(code)
        pcs_box = (cfg.pcs_per_box_override if cfg and cfg.pcs_per_box_override > 0 else (r.pack_ratio if r else 0.0)) or 0.0
        price = prices.get(code, 0.0)
        avg_pcs = sum(d["months"]) / len(d["months"]) if d["months"] else 0.0
        single_pcs = sum(single_sales.get(code, {}).get(s_, 0.0) for s_ in single_slots)
        avg_box = _round_unit(avg_pcs / pcs_box) if pcs_box > 0 else None
        single_box = _round_unit(single_pcs / pcs_box) if pcs_box > 0 else None
        avg_amt = round(avg_pcs * price, 2)
        single_amt = round(single_pcs * price, 2)
        boxes_pallet = 0
        if cfg:
            boxes_pallet = int(calc(cfg, tpls.get(cfg.template_id), r)["boxes_per_pallet"] or 0)
        size = boxes_pallet if boxes_pallet > 0 else no_pallet
        brand = d["brand"] or ""
        in_brand = brand.strip().lower() in brands
        if in_brand:
            if boxes_pallet > 0 and pcs_box > 0:
                p_avg = round(avg_pcs / pcs_box / boxes_pallet, 2)
                p_single = round(single_pcs / pcs_box / boxes_pallet, 2)
            else:
                p_avg = p_single = no_pallet
        else:
            p_avg = p_single = None
        ri += 1
        vals = [code, d["name"], brand, pcs_box or None, price or None, avg_box, avg_amt, single_box, single_amt,
                size, p_avg, p_single]
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(ri, ci, v)
            cell.border = border
            if ci in (5, 7, 9):
                cell.number_format = "#,##0"
        if size == no_pallet:
            ws.cell(ri, 10).fill = missing_fill
        if in_brand:
            for ci in (11, 12):
                ws.cell(ri, ci).fill = pallet_fill if isinstance(p_avg, float) else missing_fill
        tot_avg_amt += avg_amt
        tot_single_amt += single_amt
        tot_avg_box += avg_box or 0
        tot_single_box += single_box or 0

    ri += 1
    total_vals = {1: "Нийт", 6: tot_avg_box, 7: round(tot_avg_amt, 2), 8: tot_single_box, 9: round(tot_single_amt, 2)}
    for ci in range(1, len(cols) + 1):
        cell = ws.cell(ri, ci, total_vals.get(ci))
        cell.border = border
        cell.font = Font(bold=True)
        if ci in (7, 9):
            cell.number_format = "#,##0"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(ri - 1, 2)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    info = wb.create_sheet("Тайлбар")
    notes = [
        ("Шүүлт", f"Байршил: {tag.strip() or 'Бүх бараа'} · төлөв: {status} · хайлт: {q or '—'}"),
        ("Сарын дундаж", f"{avg_label}-р сарын борлуулалтын дундаж (агуулах + заал) — сар бүрийн хамгийн сүүлийн жилийн дата"),
        ("Тусгай сар", single_label),
        ("Хайрцаг", "Ширхэгийг хайрцаг дахь ширхэгт хувааж нэгжийн орноор тоймлосон (0.5 → дээш)"),
        ("Нэгж үнэ", "Барааны мастерын сүүлийн орлогын үнэ; дүн = борлуулалт (ширхэг) × нэгж үнэ"),
        ("Поддон хэмжээ", "Нэг поддонд багтах хайрцаг (нэг үед × үе)"),
        ("Поддоноор тооцох брендүүд", ", ".join(brand_names) or "—"),
        ("Бүрдүүлсэн", f"{datetime.now():%Y-%m-%d %H:%M} · {getattr(u, 'username', '')}"),
    ]
    for i, (k, v) in enumerate(notes, 1):
        info.cell(i, 1, k).font = Font(bold=True)
        info.cell(i, 2, v)
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    safe_tag = re.sub(r"[^0-9A-Za-zА-Яа-яӨөҮүЁё]+", "_", tag.strip())
    fname = f"{datetime.now():%Y%m%d}_borluulalt_poddon{'_' + safe_tag if tag.strip() else ''}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})
