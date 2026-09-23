"""Захиалгын түүх (Log) — нэг захиалгын бүх аудит бичлэгийг хүнд ойлгомжтой хэлбэрт
оруулж, бүх нөхцөлөөр шүүнэ.

  GET /purchase-orders/{id}/history          — {total, items, facets, sums}
  GET /purchase-orders/{id}/history/export   — Excel (ижил шүүлтээр)

Шүүлт: q (чөлөөт текст), brand, product (код/нэрийн хэсэг), user, kinds (таслалаар),
date_from/date_to (YYYY-MM-DD, серверийн локал өдөр), page/size.
Нэг захиалгад хэдэн мянган бичлэг л байдаг тул бүгдийг уншаад Python-д шүүнэ
(ix_audit_logs_parent индекстэй).
"""
from __future__ import annotations

import io
import json
from collections import Counter
from datetime import datetime, timedelta
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import and_, or_
from sqlalchemy.orm import Session

from app.api.deps import get_db, require_role
from app.models.audit_log import AuditLog
from app.models.logistics import Vehicle
from app.models.product import Product
from app.models.purchase_order import POShipment, PurchaseOrder
from app.models.user import User

router = APIRouter(prefix="/purchase-orders", tags=["purchase-order-history"])
VIEW_ROLES = ("admin", "supervisor", "manager")

STATUS_LABEL = {
    "preparing": "Бэлдэж байна", "reviewing": "Хянаж байна", "sending": "Илгээж байна",
    "loading": "Ачигдаж байна", "transit": "Замд", "arrived": "Ачаа ирсэн",
    "accounting": "Нягтлан", "confirmed": "Баталгаажсан", "received": "Орлого авагдсан",
    "cancelled": "Цуцлагдсан",
}
ACTION_LABEL = {
    "po_set_lines": "Утга шинэчилсэн",
    "po_delete_line": "Мөр устгасан",
    "po_brand_zeroed": "Брендийг бүхэлд нь 0 болгосон",
    "po_brand_advance": "Брендийн статус ахиулсан",
    "po_force_status_brand": "Брендийн статус солисон",
    "po_force_status_all": "Бүх брендийн статус солисон",
    "po_archive": "Архивласан",
    "po_unarchive": "Архиваас гаргасан",
    "po_shipment_plan_brand": "Машинд хуваарилсан",
    "po_shipment_unplan_brand": "Машинаас хассан",
    "po_set_lines_conflict": "Зөрчилтэй хадгалалт (хөндөөгүй)",
    "po_delete": "Захиалга устгасан",
    "po_erkhet_import": "Эрхэт рүү импортолсон",
}
KIND_LABEL = {
    "zeroed": "Захиалгын тоог 0 болгосон",
    "added": "Шинээр тоо оруулсан",
    "increased": "Тоо нэмсэн",
    "decreased": "Тоо хассан",
    "brand_zeroed": "Брендийг 0 болгосон",
    "line_deleted": "Мөр устгасан",
    "supplier": "Нийлүүлэгчийн тоо",
    "loaded": "Ачигдсан тоо",
    "received": "Хүлээн авсан тоо",
    "price": "Үнэ",
    "remark": "Тайлбар",
    "status": "Статус",
    "vehicle": "Машин хуваарилалт",
    "archive": "Архив",
    "conflict": "Зөрчил",
    "erkhet": "Эрхэт импорт",
    "other": "Бусад",
}
FIELD_LABEL = {
    "order_qty_box": "Захиалга (хайрцаг)", "order_qty_pcs": "Захиалга (ш)",
    "supplier_qty_box": "Нийлүүлэгч (хайрцаг)", "loaded_qty_box": "Ачигдсан (хайрцаг)",
    "received_qty_box": "Хүлээн авсан (хайрцаг)", "received_qty_extra_pcs": "Хүлээн авсан (задгай ш)",
    "unit_price": "Нэгж үнэ", "line_remark": "Тайлбар",
}
FIELD_KIND = {"supplier_qty_box": "supplier", "loaded_qty_box": "loaded", "received_qty_box": "received",
              "received_qty_extra_pcs": "received", "unit_price": "price", "line_remark": "remark"}


def _j(s):
    try:
        v = json.loads(s) if s else {}
        return v if isinstance(v, dict) else {}
    except Exception:
        return {}


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _local_offset() -> timedelta:
    return datetime.now().astimezone().utcoffset() or timedelta(0)


def _event(r: AuditLog, code_of: dict, vehicle_of: dict) -> dict:
    b, a, ex = _j(r.before_value), _j(r.after_value), _j(r.extra)
    kinds: list[str] = []
    changes: list[dict] = []
    brand = product = item_code = ""
    status = ""
    qty_delta = 0.0

    if r.action == "po_set_lines":
        brand, product, status = ex.get("brand", ""), ex.get("product_name", ""), ex.get("effective_status", "")
        item_code = code_of.get(ex.get("product_id"), "")
        for f, lbl in FIELD_LABEL.items():
            bv, av = b.get(f), a.get(f)
            if f == "line_remark":
                if (bv or "") != (av or ""):
                    changes.append({"field": f, "label": lbl, "before": bv or "", "after": av or ""})
                    kinds.append("remark")
                continue
            if f == "order_qty_pcs" or abs(_num(bv) - _num(av)) < 1e-9:
                continue
            changes.append({"field": f, "label": lbl, "before": _num(bv), "after": _num(av)})
            if f == "order_qty_box":
                qty_delta = _num(av) - _num(bv)
                if _num(bv) > 0 and _num(av) == 0:
                    kinds.append("zeroed")
                elif _num(bv) == 0 and _num(av) > 0:
                    kinds.append("added")
                elif _num(av) > _num(bv):
                    kinds.append("increased")
                else:
                    kinds.append("decreased")
            elif FIELD_KIND.get(f) and FIELD_KIND[f] not in kinds:
                kinds.append(FIELD_KIND[f])
    elif r.action == "po_delete_line":
        brand, product, status = b.get("brand", ""), b.get("product_name", ""), b.get("effective_status", "")
        item_code = code_of.get(b.get("product_id"), "")
        qty_delta = -_num(b.get("order_qty_box"))
        changes.append({"field": "order_qty_box", "label": FIELD_LABEL["order_qty_box"], "before": _num(b.get("order_qty_box")), "after": None})
        kinds.append("line_deleted")
    elif r.action == "po_brand_zeroed":
        brand = b.get("brand", "")
        qty_delta = -_num(b.get("total_order_qty_box"))
        changes.append({"field": "brand_total", "label": f"Брендийн нийт ({b.get('line_count', '')} мөр)",
                        "before": _num(b.get("total_order_qty_box")), "after": 0.0})
        kinds.append("brand_zeroed")
    elif r.action in ("po_brand_advance", "po_force_status_brand"):
        brand = b.get("brand") or a.get("brand") or ""
        changes.append({"field": "status", "label": "Статус", "before": STATUS_LABEL.get(b.get("status"), b.get("status") or ""),
                        "after": STATUS_LABEL.get(a.get("status"), a.get("status") or "")})
        kinds.append("status")
    elif r.action == "po_force_status_all":
        bc = ex.get("brand_changes") or []
        moved = [x for x in bc if x.get("from") != x.get("to")]
        changes.append({"field": "status", "label": f"Статус ({len(moved)}/{len(bc)} бренд өөрчлөгдсөн)",
                        "before": STATUS_LABEL.get(b.get("status"), b.get("status") or ""),
                        "after": STATUS_LABEL.get(a.get("status"), a.get("status") or "")})
        kinds.append("status")
    elif r.action in ("po_archive", "po_unarchive"):
        changes.append({"field": "is_archived", "label": "Архив", "before": "Тийм" if b.get("is_archived") else "Үгүй",
                        "after": "Тийм" if a.get("is_archived") else "Үгүй"})
        kinds.append("archive")
    elif r.action in ("po_shipment_plan_brand", "po_shipment_unplan_brand"):
        brand = b.get("brand") or a.get("brand") or ""
        bs, as_ = b.get("shipment_id"), a.get("shipment_id")
        changes.append({"field": "vehicle", "label": "Машин",
                        "before": vehicle_of.get(bs, f"Ачилт #{bs}") if bs else "—",
                        "after": vehicle_of.get(as_, f"Ачилт #{as_}") if as_ else "—"})
        kinds.append("vehicle")
    elif r.action == "po_erkhet_import":
        brand = a.get("brand", "")
        st = {"ok": "Амжилттай", "fail": "Алдаатай", "unknown": "Тодорхойгүй"}.get(a.get("status"), a.get("status", ""))
        changes.append({"field": "erkhet", "label": f"«{a.get('title', '')}» · {a.get('rows', 0)} мөр", "before": None,
                        "after": f"{st}" + (f" · Эрхэт #{a.get('erkhet_import_id')} · {a.get('doc_count')} баримт" if a.get("erkhet_import_id") else "")})
        if ex.get("message"):
            changes.append({"field": "erkhet_msg", "label": ex.get("message"), "before": None, "after": None})
        kinds.append("erkhet")
    elif r.action == "po_set_lines_conflict":
        n = ex.get("count", 0)
        items = ex.get("items") or []
        brands = sorted({x.get("brand", "") for x in items if x.get("brand")})
        brand = brands[0] if len(brands) == 1 else ""
        product = ", ".join(x.get("product_name", "") for x in items[:3]) + (" …" if len(items) > 3 else "")
        changes.append({"field": "conflict", "label": f"{n} мөр өөр хүн өөрчилсөн тул хадгалаагүй", "before": None, "after": None})
        kinds.append("conflict")
    else:
        kinds.append("other")
    if not kinds:
        kinds.append("other")
    return {
        "id": r.id,
        "at": (r.created_at.isoformat() + "Z") if r.created_at else None,
        "username": r.username or "",
        "role": r.role or "",
        "ip": r.ip_address or "",
        "action": r.action,
        "action_label": ACTION_LABEL.get(r.action, r.action),
        "kinds": kinds,
        "brand": brand,
        "product": product,
        "item_code": item_code,
        "status": STATUS_LABEL.get(status, status),
        "changes": changes,
        "qty_delta": round(qty_delta, 3),
    }


def _load_events(db: Session, order_id: int) -> list[dict]:
    rows = (
        db.query(AuditLog)
        .filter(or_(and_(AuditLog.parent_type == "purchase_order", AuditLog.parent_id == order_id),
                    and_(AuditLog.entity_type == "purchase_order", AuditLog.entity_id == order_id)))
        .order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
        .all()
    )
    pids = set()
    for r in rows:
        if r.action in ("po_set_lines", "po_delete_line"):
            pid = (_j(r.extra) if r.action == "po_set_lines" else _j(r.before_value)).get("product_id")
            if pid:
                pids.add(pid)
    code_of = {pid: code for pid, code in db.query(Product.id, Product.item_code).filter(Product.id.in_(pids)).all()} if pids else {}
    vehicle_of = {
        sid: f"{name} ({plate})" if name else f"Ачилт #{sid}"
        for sid, name, plate in db.query(POShipment.id, Vehicle.name, Vehicle.plate)
        .outerjoin(Vehicle, Vehicle.id == POShipment.vehicle_id)
        .filter(POShipment.purchase_order_id == order_id).all()
    }
    return [_event(r, code_of, vehicle_of) for r in rows]


def _filter(events: list[dict], q: str, brand: str, product: str, user: str, kinds: str,
            date_from: str, date_to: str) -> list[dict]:
    out = events
    if brand:
        out = [e for e in out if e["brand"] == brand]
    if user:
        out = [e for e in out if e["username"] == user]
    if product:
        p = product.strip().lower()
        out = [e for e in out if p in e["product"].lower() or p in (e["item_code"] or "").lower()]
    if kinds:
        ks = {k for k in kinds.split(",") if k}
        out = [e for e in out if ks & set(e["kinds"])]
    off = _local_offset()
    try:
        if date_from:
            lo = (datetime.strptime(date_from, "%Y-%m-%d") - off).isoformat()
            out = [e for e in out if e["at"] and e["at"][:-1] >= lo]
        if date_to:
            hi = (datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1) - off).isoformat()
            out = [e for e in out if e["at"] and e["at"][:-1] < hi]
    except ValueError:
        raise HTTPException(400, "Огноо YYYY-MM-DD байх ёстой")
    if q:
        t = q.strip().lower()
        def hay(e):
            parts = [e["username"], e["brand"], e["product"], e["item_code"], e["action_label"], e["status"]]
            parts += [KIND_LABEL.get(k, k) for k in e["kinds"]]
            parts += [f"{c['label']} {c['before']} {c['after']}" for c in e["changes"]]
            return " ".join(str(x) for x in parts if x).lower()
        out = [e for e in out if t in hay(e)]
    return out


def _facets(events: list[dict]) -> dict:
    users, brands, kinds = Counter(), Counter(), Counter()
    for e in events:
        users[e["username"]] += 1
        if e["brand"]:
            brands[e["brand"]] += 1
        for k in e["kinds"]:
            kinds[k] += 1
    return {
        "users": [{"name": n, "count": c} for n, c in users.most_common()],
        "brands": [{"name": n, "count": c} for n, c in sorted(brands.items())],
        "kinds": [{"key": k, "label": KIND_LABEL.get(k, k), "count": kinds[k]} for k in KIND_LABEL if kinds.get(k)],
    }


def _sums(events: list[dict]) -> dict:
    return {
        "added_boxes": round(sum(e["qty_delta"] for e in events if e["qty_delta"] > 0), 1),
        "removed_boxes": round(-sum(e["qty_delta"] for e in events if e["qty_delta"] < 0), 1),
    }


def _get_po(db: Session, order_id: int) -> PurchaseOrder:
    po = db.query(PurchaseOrder).filter(PurchaseOrder.id == order_id).first()
    if not po:
        # Устгагдсан захиалгын түүхийг ч харж болно (аудит үлддэг)
        if not db.query(AuditLog.id).filter(AuditLog.parent_type == "purchase_order", AuditLog.parent_id == order_id).first():
            raise HTTPException(404, "Захиалга олдсонгүй")
    return po


@router.get("/{order_id}/history")
def order_history(
    order_id: int,
    q: str = "", brand: str = "", product: str = "", user: str = "", kinds: str = "",
    date_from: str = "", date_to: str = "",
    page: int = Query(1, ge=1), size: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db),
    _: User = Depends(require_role(*VIEW_ROLES)),
):
    _get_po(db, order_id)
    events = _load_events(db, order_id)
    filtered = _filter(events, q, brand, product, user, kinds, date_from, date_to)
    start = (page - 1) * size
    return {
        "order_id": order_id,
        "total_all": len(events),
        "total": len(filtered),
        "page": page, "size": size,
        "items": filtered[start:start + size],
        "facets": _facets(events),
        "sums": _sums(filtered),
    }


@router.get("/{order_id}/history/export")
def order_history_export(
    order_id: int,
    q: str = "", brand: str = "", product: str = "", user: str = "", kinds: str = "",
    date_from: str = "", date_to: str = "",
    db: Session = Depends(get_db),
    _: User = Depends(require_role(*VIEW_ROLES)),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    po = _get_po(db, order_id)
    events = _filter(_load_events(db, order_id), q, brand, product, user, kinds, date_from, date_to)
    off = _local_offset()
    wb = Workbook()
    ws = wb.active
    ws.title = "Түүх"
    heads = ["Огноо", "Хэрэглэгч", "Үйлдэл", "Төрөл", "Бренд", "Код", "Бараа", "Талбар", "Өмнө", "Дараа", "Статус", "IP"]
    ws.append(heads)
    thin = Side(style="thin", color="BBBBBB")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", start_color="DDEBF7")
        c.border = bd
    for e in events:
        at = datetime.fromisoformat(e["at"][:-1]) + off if e["at"] else None
        base = [at.strftime("%Y-%m-%d %H:%M:%S") if at else "", e["username"], e["action_label"],
                ", ".join(KIND_LABEL.get(k, k) for k in e["kinds"]), e["brand"], e["item_code"], e["product"]]
        chs = e["changes"] or [{"label": "", "before": "", "after": ""}]
        for ch in chs:
            ws.append(base + [ch["label"], ch["before"] if ch["before"] is not None else "",
                              ch["after"] if ch["after"] is not None else "", e["status"], e["ip"]])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.border = bd
            c.alignment = Alignment(vertical="top")
    for col, w in zip("ABCDEFGHIJKL", (19, 16, 24, 22, 26, 10, 40, 22, 12, 12, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    date = po.order_date.isoformat() if po else ""
    fname = f"Захиалга_{order_id}_{date}_түүх.xlsx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}"})
