"""Улайлтын тайлан — Эрхэтэд сөрөг үлдэгдэлтэй бараа (app/services/erkhet_red_report).

  GET  /red-report            — сүүлийн үр дүн + тохиргоо + Эрхэтийн байршлууд + төлөв
  PUT  /red-report/settings   — {locations, interval_min, start_hour, end_hour}
  POST /red-report/run        — гараар шинэчлэх (5 минутын завсартай)
  GET  /red-report/xlsx       — «Улайлт тайлан» Excel
"""
from __future__ import annotations

import io
import threading
import time
from datetime import datetime, timedelta
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from app.api.deps import require_role
from app.models.user import User
from app.services import erkhet_red_report as svc

router = APIRouter(prefix="/red-report", tags=["red-report"])
VIEW = ("admin", "supervisor", "manager")
_locs: dict = {"at": 0.0, "v": []}


def _all_locations() -> list[dict]:
    if not _locs["v"] or time.monotonic() - _locs["at"] > 3600:
        try:
            _locs.update(at=time.monotonic(), v=svc.locations())
        except Exception:
            pass
    return _locs["v"]


@router.get("")
def get_report(u: User = Depends(require_role(*VIEW))):
    return {"result": svc.load_result(), "settings": svc.load_settings(), "status": svc.status(),
            "all_locations": _all_locations()}


class SettingsIn(BaseModel):
    locations: Optional[list[str]] = None
    interval_min: Optional[int] = None
    start_hour: Optional[int] = None
    end_hour: Optional[int] = None


@router.put("/settings")
def put_settings(body: SettingsIn, u: User = Depends(require_role(*VIEW))):
    return {"settings": svc.save_settings(body.model_dump(exclude_none=True))}


@router.post("/run")
def run(u: User = Depends(require_role(*VIEW))):
    st = svc.status()
    if st["running"]:
        return {"started": False, "status": st}
    if st["next_manual_in"]:
        raise HTTPException(429, f"Эрхэтэд ачаалал үүсгэхгүйн тулд {st['next_manual_in']} секундын дараа дахин оролдоно уу.")
    threading.Thread(target=svc.run, args=(True,), daemon=True, name="red-report").start()
    return {"started": True}


@router.get("/xlsx")
def xlsx(u: User = Depends(require_role(*VIEW))):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    r = svc.load_result()
    if not r:
        raise HTTPException(404, "Улайлтын тайлан хараахан гараагүй.")
    at = datetime.fromisoformat(r["checked_at"].rstrip("Z")) + timedelta(hours=8)
    wb = Workbook()
    ws = wb.active
    ws.title = "Улайлт тайлан"
    ws.append(["Улайлт тайлан"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Эрхэтийн үлдэгдэл {at:%Y-%m-%d %H:%M}-ийн байдлаар · байршил: "
               + ", ".join(f"{x['code']} {x['name']}" for x in r["locations"])])
    ws.append([])
    ws.append(["Байршил", "Код", "Нэр", "Эхний үлдэгдэл", "Орлого", "Зарлага", "Эцсийн үлдэгдэл", "Дүн", "Нэгж өртөг"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", start_color="F3F4F6")
    red = Font(color="DC2626", bold=True)
    cur = None
    for it in r["items"]:
        if it["location"] != cur:
            cur = it["location"]
            n = r["by_location"].get(cur, 0)
            ws.append([f"{cur} {it['location_name']} — {n} бараа"])
            ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.append([it["location"], it["code"], it["name"], it["begin_qty"], it["in_qty"], it["out_qty"], it["qty"],
                   it["amount"], it["unit_cost"]])
        ws.cell(ws.max_row, 7).font = red
        for col in (4, 5, 6, 7, 8, 9):
            ws.cell(ws.max_row, col).number_format = "#,##0.##;-#,##0.##"
    ws.append([])
    ws.append([f"Нийт: {r['count']} бараа"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for col, w in zip("ABCDEFGHI", (10, 12, 44, 14, 10, 10, 15, 13, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "4:4"
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"Улайлт_тайлан_{at:%Y-%m-%d_%H%M}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fn)}"})
