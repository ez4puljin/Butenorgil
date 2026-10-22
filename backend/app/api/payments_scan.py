"""paymentsData илрүүлэгч ба засагч — Erkhet sync шалгалтын дэд хэсэг (app/services/payments_scan).

  GET  /erkhet-sync/payments/settings         — зөвшөөрөгдсөн төлбөрийн төрлүүд (PUT — admin)
  POST /erkhet-sync/payments-scan             — {boardIds?, pipelineIds?, startDate, endDate, categories?} → run_id
  GET  /erkhet-sync/payments-scan/runs        — сүүлийн скан/засварууд (?kind=scan|auto|fix)
  GET  /erkhet-sync/payments-scan/runs/{id}   — явц + deal бүрийн үр дүн (засвар бол actions)
  GET  /erkhet-sync/payments-scan/runs/{id}/xlsx
  POST /erkhet-sync/payments-fix/preview      — dry-run: erxes-т юу ч бичихгүй → token
  POST /erkhet-sync/payments-fix              — {token, only_first, confirm} (admin) → run_id
  GET  /erkhet-sync/payments-fix/actions      — засварын audit (хуучин paymentsData — rollback)
  POST /erkhet-sync/payments-fix/rollback     — {action_id, confirm} (admin) — paymentsData-г хуучнаар нь

Хатуу дүрэм: автоматаар засахгүй, sync хийхгүй. dealsEditProductData ашиглахгүй.
"""
from __future__ import annotations

import io
import json
import threading
from datetime import date, datetime, timedelta
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.api.deps import get_db, require_role
from app.core.audit import audit
from app.models.payments_scan import PaymentsFixAction, PaymentsScanItem, PaymentsScanRun
from app.models.user import User
from app.services import erkhet_deal_check as K
from app.services import payments_scan as svc

router = APIRouter(prefix="/erkhet-sync", tags=["erkhet-sync"])
VIEW = ("admin", "supervisor", "manager")
FIX = ("admin",)                    # засах/sync — зөвхөн админ


class TypesIn(BaseModel):
    types: list[str]


@router.get("/payments/settings")
def get_settings(db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    return {"allowed": svc.load_allowed(db), "issues": {k: list(v) for k, v in svc.ISSUE_INFO.items()}}


@router.put("/payments/settings")
def put_settings(body: TypesIn, db: Session = Depends(get_db), u: User = Depends(require_role("admin"))):
    return {"allowed": svc.save_allowed(db, body.types, u.username)}


class ScanIn(BaseModel):
    boardIds: list[str] = []
    pipelineIds: list[str] = []
    categories: list[str] = []
    startDate: str
    endDate: str


def _busy(db: Session, kinds: tuple[str, ...]) -> PaymentsScanRun | None:
    since = datetime.utcnow() - timedelta(hours=2)
    return (db.query(PaymentsScanRun).filter(PaymentsScanRun.kind.in_(kinds), PaymentsScanRun.status == "running",
                                             PaymentsScanRun.created_at >= since).first())


@router.post("/payments-scan")
def start_scan(body: ScanIn, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    try:
        a, b = date.fromisoformat(body.startDate), date.fromisoformat(body.endDate)
    except ValueError:
        raise HTTPException(400, "Огноо буруу (YYYY-MM-DD).")
    if a > b or (b - a).days > 92:
        raise HTTPException(400, "Огнооны хүрээ 1–92 хоног байна.")
    if set(body.categories) - set(K.CATEGORIES):
        raise HTTPException(400, "Stage-ийн ангилал буруу.")
    busy = _busy(db, ("scan",))
    if busy:
        return {"run_id": busy.id, "reused": True}
    params = {**body.model_dump(), "dateField": "createdAt"}
    r = PaymentsScanRun(kind="scan", status="running", params=json.dumps(params), created_by=u.username)
    db.add(r)
    db.commit()
    threading.Thread(target=svc.run_scan, args=(r.id, params), daemon=True, name=f"payments-scan-{r.id}").start()
    return {"run_id": r.id, "reused": False}


def _run_out(r: PaymentsScanRun) -> dict:
    return {"id": r.id, "kind": r.kind, "status": r.status, "params": json.loads(r.params or "{}"),
            "progress": r.progress, "done_n": r.done_n, "total_n": r.total_n,
            "summary": json.loads(r.summary or "{}"), "error": r.error, "created_by": r.created_by,
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "finished_at": r.finished_at.isoformat() if r.finished_at else None}


def _item_out(x: PaymentsScanItem) -> dict:
    return {"deal_id": x.deal_id, "number": x.number, "name": x.name, "pipeline_id": x.pipeline_id,
            "stage_id": x.stage_id, "category": x.category, "erxes_amount": x.erxes_amount,
            "erkhet_amount": x.erkhet_amount, "amount_diff": x.amount_diff, "erkhet_payment": x.erkhet_payment,
            "last_sync_undefined": x.last_sync_undefined, "payments_bad": x.payments_bad, "severity": x.severity,
            "issues": json.loads(x.issues or "[]"), "data": json.loads(x.data or "{}")}


def _action_out(a: PaymentsFixAction) -> dict:
    return {"id": a.id, "run_id": a.run_id, "deal_id": a.deal_id, "number": a.number, "action": a.action,
            "payments_before": a.payments_before, "payments_after": a.payments_after,
            "snapshot_before": json.loads(a.snapshot_before or "{}"), "snapshot_after": json.loads(a.snapshot_after or "{}"),
            "config_title": a.config_title, "date_type": a.date_type, "erkhet_before": a.erkhet_before,
            "erkhet_after": a.erkhet_after, "result": a.result, "message": a.message, "user": a.user_name,
            "at": a.created_at.isoformat() if a.created_at else None}


@router.get("/payments-scan/runs")
def list_runs(kind: str = "", limit: int = 20, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    q = db.query(PaymentsScanRun)
    if kind:
        q = q.filter(PaymentsScanRun.kind.in_(kind.split(",")))
    return [_run_out(r) for r in q.order_by(PaymentsScanRun.id.desc()).limit(min(limit, 100))]


@router.get("/payments-scan/runs/{run_id}")
def get_run(run_id: int, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    r = db.get(PaymentsScanRun, run_id)
    if not r:
        raise HTTPException(404, "Олдсонгүй.")
    out = _run_out(r)
    if r.status == "done" and r.kind in ("scan", "auto"):
        out["items"] = [_item_out(x) for x in db.query(PaymentsScanItem).filter(PaymentsScanItem.run_id == run_id)
                        .order_by(PaymentsScanItem.id)]
    if r.kind == "fix":
        out["actions"] = [_action_out(a) for a in db.query(PaymentsFixAction).filter(PaymentsFixAction.run_id == run_id)
                          .order_by(PaymentsFixAction.id)]
    return out


@router.get("/payments-scan/runs/{run_id}/xlsx")
def export_xlsx(run_id: int, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    r = db.get(PaymentsScanRun, run_id)
    if not r or r.status != "done":
        raise HTTPException(404, "Дууссан скан олдсонгүй.")
    p = json.loads(r.params or "{}")
    wb = Workbook()
    ws = wb.active
    ws.title = "paymentsData"
    ws.append(["Deal №", "Нэр", "Харилцагч", "Board", "Pipeline", "Stage", "Ангилал", "Erxes дүн", "Erkhet дүн",
               "Зөрүү", "Erkhet төлбөр", "paymentsData", "Асуудал", "Ноцтой", "Сүүлийн sync (УБ)", "Үүссэн (УБ)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    fills = {"error": "FDE2E2", "warn": "FEF3C7", "info": "E0F2FE"}
    for x in db.query(PaymentsScanItem).filter(PaymentsScanItem.run_id == run_id).order_by(PaymentsScanItem.id):
        d = json.loads(x.data or "{}")
        iss = json.loads(x.issues or "[]")
        ws.append([x.number, x.name, d.get("customer", ""), d.get("board", ""), d.get("pipeline", ""), d.get("stage", ""),
                   x.category, x.erxes_amount, x.erkhet_amount, x.amount_diff if x.erkhet_amount is not None else None,
                   x.erkhet_payment, json.dumps(d.get("paymentsData"), ensure_ascii=False),
                   ", ".join(svc.ISSUE_INFO.get(c, ("", c))[1] for c in iss), x.severity,
                   K.ub_date(d.get("last_sync_at")), K.ub_date(d.get("createdAt"))])
        for col in (8, 9, 10):
            ws.cell(ws.max_row, col).number_format = "#,##0;[Red]-#,##0"
        if x.severity in fills:
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", start_color=fills[x.severity])
    for col, w in zip("ABCDEFGHIJKLMNOP", (16, 30, 24, 10, 22, 22, 10, 13, 13, 11, 13, 34, 50, 8, 12, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"paymentsData_{p.get('startDate', '')}_{p.get('endDate', '')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fn)}"})


class PreviewIn(BaseModel):
    deal_ids: list[str]
    mode: str = "empty"                 # empty → {} | other → {"other": {amount: tickUsed нийлбэр}}


@router.post("/payments-fix/preview")
def fix_preview(body: PreviewIn, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    """Dry-run — erxes-т юу ч бичихгүй."""
    ids = list(dict.fromkeys(body.deal_ids))
    if not ids:
        raise HTTPException(400, "Deal сонгоогүй байна.")
    if len(ids) > 200:
        raise HTTPException(400, "Нэг удаад 200-аас ихгүй deal.")
    if body.mode not in ("empty", "other"):
        raise HTTPException(400, "Засах арга буруу.")
    plan = svc.build_plan(ids, body.mode, svc.load_allowed(db))
    return {"token": svc.store_plan(plan, body.mode), "mode": body.mode, "plan": plan,
            "counts": {k: sum(1 for x in plan if x["action"] == k) for k in ("fix_sync", "fix_only", "sync_only", "skip")}}


class FixIn(BaseModel):
    token: str
    only_first: bool = True             # эхлээд 1 deal — үр дүнг харсны дараа л олноор
    confirm: bool = False


@router.post("/payments-fix")
def fix_run(body: FixIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_role(*FIX))):
    if not body.confirm:
        raise HTTPException(400, "Dry-run-ийг баталгаажуулна уу (confirm).")
    p = svc.take_plan(body.token)
    if not p:
        raise HTTPException(400, "Dry-run хуучирсан (15 мин) — дахин dry-run хийнэ үү.")
    todo = [x for x in p["plan"] if x["action"] != "skip"]
    if not todo:
        raise HTTPException(400, "Хийх зүйл алга.")
    if body.only_first:
        todo = todo[:1]
    busy = K.lock_deals([x["deal_id"] for x in todo])
    if busy:
        raise HTTPException(409, f"{len(busy)} deal дээр өөр засвар/sync ажиллаж байна.")
    if _busy(db, ("fix",)):
        K.unlock_deals([x["deal_id"] for x in todo])
        raise HTTPException(409, "Өөр засвар ажиллаж байна.")
    r = PaymentsScanRun(kind="fix", status="running", total_n=len(todo), created_by=u.username,
                        params=json.dumps({"mode": p["mode"], "only_first": body.only_first,
                                           "deals": [x["number"] for x in todo]}, ensure_ascii=False))
    db.add(r)
    db.commit()
    audit(db, request, u, action="payments_fix", entity_type="erkhet_sync",
          extra={"run_id": r.id, "mode": p["mode"], "deals": [x["number"] for x in todo]}, autocommit=True)
    if body.only_first:
        # Үлдсэнийг ижил dry-run-аар үргэлжлүүлэхэд token хүчинтэй хэвээр — эхнийхийг хасна
        p["plan"] = [x for x in p["plan"] if x["deal_id"] != todo[0]["deal_id"]]
    else:
        svc._PLANS.pop(body.token, None)
    threading.Thread(target=svc.run_fix, args=(r.id, todo, u.username), daemon=True, name=f"payments-fix-{r.id}").start()
    return {"run_id": r.id, "n": len(todo), "remaining": sum(1 for x in p["plan"] if x["action"] != "skip")
            if body.only_first else 0}


@router.get("/payments-fix/actions")
def fix_actions(limit: int = 200, db: Session = Depends(get_db), u: User = Depends(require_role(*VIEW))):
    return [_action_out(a) for a in db.query(PaymentsFixAction).order_by(PaymentsFixAction.id.desc()).limit(min(limit, 1000))]


class RollbackIn(BaseModel):
    action_id: int
    confirm: bool = False


@router.post("/payments-fix/rollback")
def rollback(body: RollbackIn, request: Request, db: Session = Depends(get_db), u: User = Depends(require_role(*FIX))):
    """paymentsData-г засахаас өмнөх утгаар нь буцаана (sync хийхгүй)."""
    a = db.get(PaymentsFixAction, body.action_id)
    if not a or not a.payments_before:
        raise HTTPException(404, "Буцаах paymentsData алга.")
    if not body.confirm:
        return {"deal": a.number, "will_set": json.loads(a.payments_before)}
    before = svc.fetch_full_detail(a.deal_id)
    K._gql('mutation($id: String!, $p: JSON) { dealsEdit(_id: $id, paymentsData: $p) { _id } }',
           {"id": a.deal_id, "p": json.loads(a.payments_before)}, tries=1)
    after = svc.fetch_full_detail(a.deal_id)
    sb, sa = svc._snapshot(before), svc._snapshot(after)
    changed = [k for k in ("rows", "products", "amount", "stageId", "name") if sb[k] != sa[k]]
    audit(db, request, u, action="payments_fix_rollback", entity_type="erkhet_sync",
          extra={"action_id": a.id, "deal": a.number, "changed": changed}, autocommit=True)
    if changed:
        raise HTTPException(500, f"Анхаар: {', '.join(changed)} өөрчлөгдсөн байна — шалгана уу.")
    return {"ok": True, "deal": a.number, "paymentsData": K._j(after.get("paymentsData"))}
