"""Тооллогын хуудас → Эрхэтийн импорт Excel: илүүдэл = «Бараа материалын орлого» (inv_income, 23 багана),
дутагдал = «Бараа материалын байнгын борлуулалт» (inv_csale, 25 багана — хэрэглэгчийн өгсөн Эрхэтийн
толгой файл import_sample.xls, 2026-10-31).

Тоо = тооллогын харьцуулалтын (засварласан) Зөрүү: Тооллого − Програм.
Харилцагч 60018 «Тооллого», Харьцсан данс 313011 «Тооллогын зөрүү», Данс* 150101 (тохиргоонд хадгална).

Дутагдал (байнгын борлуулалт): нэгж үнэ = ЗАРАХ ҮНЭ, НӨАТ/НХАТ-гүй — Эрхэтэд гараар хийсэн тооллогын
байнгын борлуулалтууд (харилцагч «Тооллого», дт:313011, кт:510101) яг ингэж бүртгэгдсэн (үнэ/зарах = 1.0,
кредит = бүтэн дүн); өртгийг Эрхэт өөрөө бодно.
Илүүдэл (орлого): нэгж үнэ = ӨРТӨГ — орлого жигнэсэн дундаж өртгийг өөрчилнө. Тухайн байршлын тухайн
өдрийн «Нэгж өртөг»-ийг Эрхэтийн «Бараа материал /Өртгөөр/» тайлангаас татна (нэг байршил ~1-2 мин —
арын ажлаар, диск дээр кэштэй); хэвийн биш бол агуулахын өртөг → сүүлийн авсан үнэ → зарах үнэ (шалгах).
Эрхэтийн импортын дүрэм: *-тэй баганыг мөр бүрт, бусдыг баримтын эхний мөрөнд; тийм/үгүй = 1/0.
"""
from __future__ import annotations

import io
import json
import threading
import time
import traceback
from datetime import date, datetime
from pathlib import Path

from app.services.conversion import INCOME_HEADERS

# Эрхэтийн «Бараа материалын байнгын борлуулалт» (inv_csale) импортын толгой — яг энэ дарааллаар
CSALE_HEADERS = ["Баримтын дугаар", "Огноо", "Гүйлгээний утга", "Харилцагч", "Харьцсан данс",
                 "Харьцсан ялгаатай харилцагч", "Борлуулагч", "Баримт-НӨАТ-тэй эсэх", "НӨАТ автоматаар тооцох эсэх",
                 "НӨАТ-н дүн", "НӨАТ-н үзүүлэлт", "НХАТ-тай эсэх", "НХАТ автоматаар бодох эсэх", "НХАТ ийн дүн",
                 "Бараа материал*", "Барааны данс*", "Барааны байршил*", "Тоо хэмжээ*", "Нэгж үнэ*", "Нийт дүн*",
                 "Мөр-НӨАТ тооцох эсэх*", "Мөр-НХАТ тооцох эсэх*", "НХАТ үзүүлэлт*", "Хөнгөлөгдсөн дүн*",
                 "Хариуцагч РД"]

SETTINGS_FILE = Path("app/data/count_erkhet_settings.json")
CACHE_DIR = Path("app/data/cache/count_cost")
DEFAULTS = {"customer": "60018", "related_account": "313011", "account": "150101"}
COST_REPORT = "inventory/report/generate/"           # «Бараа материал /Өртгөөр/»
COST_TTL = 6 * 3600                                    # сек — үүнээс хуучин өртгийг дахин татна
EPS = 1e-6

# Тооллогын агуулах → Эрхэтийн байршлын код (хэрэглэгч солиж болно)
WAREHOUSE_LOCATION = {
    "wholesale": "01", "drink_alcohol": "02", "hall": "03", "kharkhorin_hall": "04",
    "kharkhorin_wh": "05", "alcohol_hall": "08", "retail": "11", "contract": "12",
}
# Эрхэт хүрэхгүй үед (2026-10-31-ний жагсаалт)
FALLBACK_LOCATIONS = [
    ("01", "Бөөний агуулах"), ("02", "Ус ундаа архи пиво"), ("03", "Заал"), ("04", "Хархорин заал"),
    ("05", "Хархорин агуулах"), ("08", "Заалны архи"), ("10", "Агуулахын Хөргүүр"),
    ("11", "Жижиглэн барааны агуулах"), ("12", "Гэрээт компаний агуулах"), ("13", "Гэр ахуйн барааны агуулах"),
    ("14", "Discrepancy Control"), ("17", "Аяны барааны гадна лангуу"),
]
KIND_LABEL = {"income": "Бараа материалын орлого", "expense": "Бараа материалын байнгын борлуулалт"}
KIND_SLUG = {"income": "orlogo", "expense": "zarlaga"}


def load_settings() -> dict:
    try:
        s = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        s = {}
    return {**DEFAULTS, **{k: str(v).strip() for k, v in s.items() if k in DEFAULTS and str(v or "").strip()}}


def save_settings(s: dict) -> dict:
    cur = load_settings()
    cur.update({k: str(v).strip() for k, v in s.items() if k in DEFAULTS and str(v or "").strip()})
    SETTINGS_FILE.parent.mkdir(parents=True, exist_ok=True)
    SETTINGS_FILE.write_text(json.dumps(cur, ensure_ascii=False, indent=1), encoding="utf-8")
    return cur


# ── Эрхэтийн байршил ───────────────────────────────────────────────────────────

_loc_cache: dict = {"at": 0.0, "rows": None}


def _client():
    """Тусдаа session — 2 минутын тайлан бусад Эрхэтийн дуудлагыг (дундын клиентийн lock) хаахгүй."""
    from app.services.erkhet_client import ErkhetClient
    return ErkhetClient(timeout=900)


def locations() -> list[dict]:
    """[{code, name, id}] — Эрхэтийн тайлангийн формоос (1 цаг кэш); хүрэхгүй бол тогтмол жагсаалт (id-гүй)."""
    if _loc_cache["rows"] and time.time() - _loc_cache["at"] < 3600:
        return _loc_cache["rows"]
    try:
        from app.services.erkhet_client import get_client
        opts = get_client().form_options(COST_REPORT).get("inv_location", [])
        rows = []
        for val, label in opts:
            code, _, name = label.partition(" - ")
            rows.append({"code": code.strip(), "name": name.strip(), "id": val})
        if rows:
            _loc_cache.update(at=time.time(), rows=rows)
            return rows
    except Exception as e:  # noqa: BLE001
        print(f"[count_erkhet] байршлын жагсаалт: {e}")
    return [{"code": c, "name": n, "id": ""} for c, n in FALLBACK_LOCATIONS]


def location_name(code: str) -> str:
    return next((l["name"] for l in locations() if l["code"] == code), "")


# ── Нэгж өртөг (Эрхэт, байршил × өдөр) ─────────────────────────────────────────

def _f(v) -> float:
    try:
        return float(str(v).replace(",", "").replace(" ", "") or 0)
    except ValueError:
        return 0.0


def parse_cost_rows(rows: list[dict]) -> dict[str, float]:
    """«Бараа материал /Өртгөөр/» (Данс-Байршил) мөрүүд → {код: нэгж өртөг > 0}.
    Багана: 0 код · 1 нэр · 2-3 эхний · 4-5 орлого · 6-7 зарлага · 8 эцсийн тоо · 9 эцсийн дүн · 10 нэгж өртөг.
    Данс/байршлын толгой мөр (тоо хоосон) алгасна; өртөг 0/сөрөг бол эцсийн дүн ÷ тоо, тэр нь ч болохгүй бол алгасна."""
    out: dict[str, float] = {}
    for r in rows:
        vals = list(r.values())
        if len(vals) < 11:
            continue
        code = str(vals[0] or "").strip()
        if not code or all(str(x).strip() == "" for x in vals[2:11]):
            continue
        cost = _f(vals[10])
        if cost <= 0:
            qty, amt = _f(vals[8]), _f(vals[9])
            cost = amt / qty if qty > 0 and amt > 0 else 0.0
        if cost > 0:
            out[code] = round(cost, 2)
    return out


def fetch_costs(loc_code: str, day: date) -> dict[str, float]:
    loc = next((l for l in locations() if l["code"] == loc_code and l["id"]), None)
    if not loc:
        raise RuntimeError(f"Эрхэтэд «{loc_code}» байршил олдсонгүй.")
    c = _client()
    account = c.find_option(COST_REPORT, "account", "150101")
    ds = day.strftime("%m/%d/%Y")
    rows = c.report_table(COST_REPORT, {
        "account": account, "inv_location": [loc["id"]], "get_kind": "account_location",
        "get_tr_kind": "all", "fraction": "2", "begin_date": ds, "end_date": ds,
    })
    costs = parse_cost_rows(rows)
    if not costs:
        raise RuntimeError(f"Эрхэтийн өртгийн тайлан хоосон ирлээ ({loc_code}, {day}).")
    return costs


def _cache_path(loc_code: str, day: date) -> Path:
    return CACHE_DIR / f"{loc_code}_{day.isoformat()}.json"


def cached_costs(loc_code: str, day: date) -> dict | None:
    """{"costs", "fetched_at" (UTC ISO), "fresh"} эсвэл None."""
    try:
        d = json.loads(_cache_path(loc_code, day).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    age = (datetime.utcnow() - datetime.fromisoformat(d["fetched_at"])).total_seconds()
    return {**d, "fresh": age < COST_TTL}


_jobs: dict[tuple, dict] = {}
_job_lock = threading.Lock()


def job_state(loc_code: str, day: date) -> dict:
    """{"status": none|running|done|error, "fetched_at", "error", "n", "fresh"}"""
    j = _jobs.get((loc_code, day))
    if j and j["status"] == "running":
        return {"status": "running", "started_at": j["started_at"]}
    c = cached_costs(loc_code, day)
    if j and j["status"] == "error" and not (c and c["fetched_at"] > j["started_at"]):
        return {"status": "error", "error": j["error"], "started_at": j["started_at"],
                **({"fetched_at": c["fetched_at"], "n": len(c["costs"]), "fresh": c["fresh"]} if c else {})}
    if c:
        return {"status": "done", "fetched_at": c["fetched_at"], "n": len(c["costs"]), "fresh": c["fresh"]}
    return {"status": "none"}


def start_cost_job(loc_code: str, day: date, force: bool = False) -> dict:
    """Арын ажлаар өртөг татна. Нэг удаад нэг тайлан (Эрхэтийг ачааллахгүй)."""
    with _job_lock:
        if any(j["status"] == "running" for j in _jobs.values()):
            run = next(k for k, j in _jobs.items() if j["status"] == "running")
            if run == (loc_code, day):
                return job_state(loc_code, day)
            raise RuntimeError(f"Өөр өртөг ({run[0]}, {run[1]}) татаж байна — түр хүлээгээд дахин оролдоно уу.")
        c = cached_costs(loc_code, day)
        if c and c["fresh"] and not force:
            return job_state(loc_code, day)
        _jobs[(loc_code, day)] = {"status": "running", "started_at": datetime.utcnow().isoformat(), "error": ""}

    def run():
        j = _jobs[(loc_code, day)]
        try:
            costs = fetch_costs(loc_code, day)
            CACHE_DIR.mkdir(parents=True, exist_ok=True)
            _cache_path(loc_code, day).write_text(json.dumps(
                {"location": loc_code, "date": day.isoformat(), "fetched_at": datetime.utcnow().isoformat(),
                 "costs": costs}, ensure_ascii=False), encoding="utf-8")
            j["status"] = "done"
        except Exception as e:  # noqa: BLE001
            traceback.print_exc()
            j.update(status="error", error=str(e)[:300] or e.__class__.__name__)

    threading.Thread(target=run, daemon=True, name=f"count-cost-{loc_code}-{day}").start()
    return job_state(loc_code, day)


# ── Мөр сонгох, үнэ ─────────────────────────────────────────────────────────────

def _num(v):
    """Бүхэл бол int — Excel-д '24.0' биш '24'."""
    f = float(v)
    return int(round(f)) if abs(f - round(f)) < EPS else round(f, 3)


# Өртөг зарах үнийн энэ хүрээнд байвал хүлээн авна. 2026-10-31-ний 03 Заалны тооллогоор: байршлын үлдэгдэл
# бага үед Эрхэтийн байршлын өртөг гажуудсан (3,500₮-ийн бараа 934,237₮, 4,540₮-ийнх 2₮), «сүүлийн авсан
# үнэ» заримд нь хайрцгийн/нийт дүн. Эрүүл өртөгтэй барааны өртөг ÷ зарах үнэ медиан 0.91.
SANE_LO, SANE_HI = 0.3, 1.2
PRICE_SRC_LABEL = {"erkhet": "Эрхэт — байршлын өртөг", "warehouse": "Эрхэт — үндсэн агуулахын өртөг",
                   "last": "Сүүлийн авсан үнэ", "sell": "Зарах үнэ (тооллогын хуудас)"}


def warehouse_costs() -> dict[str, float]:
    """Шөнө бүр татдаг «Бүх агуулахын үлдэгдэл» (01·02·11·12, өчигдрийн эцсийн) → {код: тухайн барааны
    хамгийн их үлдэгдэлтэй байршлын нэгж өртөг}. Заал руу бараа агуулахаас өртгөөрөө шилждэг тул нөөц эх сурвалж."""
    from python_calamine import CalamineWorkbook
    from app.services.erkhet_sync import UPLOAD_DIR as BAL_DIR
    p = BAL_DIR / "warehouse.xls"
    if not p.exists():
        return {}
    try:
        wb = CalamineWorkbook.from_path(str(p))
        rows = wb.get_sheet_by_name(wb.sheet_names[0]).to_python()
    except Exception as e:  # noqa: BLE001
        print(f"[count_erkhet] warehouse.xls: {e}")
        return {}
    best: dict[str, tuple[float, float]] = {}
    for r in rows[2:]:
        if len(r) < 11 or r[10] in ("", None):
            continue
        code = str(r[0]).strip()
        code = code[:-2] if code.endswith(".0") else code
        q, a, c = _f(r[8]), _f(r[9]), _f(r[10])
        if c <= 0 and q > 0 and a > 0:
            c = a / q
        if c > 0 and (code not in best or q > best[code][0]):
            best[code] = (q, c)
    return {k: v[1] for k, v in best.items()}


def _sane(cost: float, sell: float) -> bool:
    return cost > 0 and (not sell or SANE_LO * sell <= cost <= SANE_HI * sell)


def priced_rows(db, compare_rows: list[dict], costs: dict[str, float], mode: str = "cost",
                lpp: dict[str, float] | None = None, wh: dict[str, float] | None = None) -> list[dict]:
    """Зөрүүтэй мөр бүрт нэгж үнэ.
    Дутагдал (байнгын борлуулалт): үргэлж зарах үнэ — өртгийг Эрхэт өөрөө бодно.
    Илүүдэл (орлого) mode="cost": Эрхэтийн байршлын өртөг → үндсэн агуулахын өртөг → сүүлийн авсан үнэ (тус бүр
    зарах үнийн 30–120%-д багтвал) → эс бөгөөс зарах үнэ (шалгах мөр). mode="sell": зарах үнэ."""
    from app.models.product import Product
    rows = [r for r in compare_rows if abs(r["diff"]) > EPS]
    if lpp is None and mode == "cost":
        codes, lpp = [r["code"] for r in rows], {}
        for i in range(0, len(codes), 900):
            for code, price in (db.query(Product.item_code, Product.last_purchase_price)
                                  .filter(Product.item_code.in_(codes[i:i + 900])).all()):
                lpp[code] = max(lpp.get(code, 0.0), float(price or 0))
    lpp = lpp or {}
    if wh is None:
        wh = warehouse_costs() if mode == "cost" and any(r["diff"] > 0 for r in rows) else {}
    out = []
    for r in rows:
        sell = float(r.get("price") or 0)
        kind = "income" if r["diff"] > 0 else "expense"
        cands = [("erkhet", costs.get(r["code"], 0.0)), ("warehouse", wh.get(r["code"], 0.0)),
                 ("last", lpp.get(r["code"], 0.0))]
        if mode == "sell" or kind == "expense":
            cands = []
        src, price = next(((s, c) for s, c in cands if _sane(c, sell)), ("sell", sell) if sell > 0 else ("none", 0.0))
        out.append({"code": r["code"], "name": r["name"], "brand": r.get("brand", ""), "diff": r["diff"],
                    "qty": abs(r["diff"]), "price": round(price, 2), "price_src": src, "sell": sell,
                    "review": kind == "income" and mode == "cost" and src in ("sell", "none"), "kind": kind})
    return out


def _brief(r: dict) -> dict:
    return {"code": r["code"], "name": r["name"], "qty": _num(r["qty"]), "price": r["price"], "sell": r["sell"],
            "src": r["price_src"]}


def preview(rows: list[dict]) -> dict:
    out = {}
    for kind in ("income", "expense"):
        sel = [r for r in rows if r["kind"] == kind]
        ok = [r for r in sel if r["price"] > 0]
        out[kind] = {
            "rows": len(ok), "pieces": _num(sum(r["qty"] for r in ok)) if ok else 0,
            "amount": round(sum(r["qty"] * r["price"] for r in ok), 2),
            "by_src": {s: sum(1 for r in ok if r["price_src"] == s) for s in PRICE_SRC_LABEL},
            "review": [_brief(r) for r in ok if r["review"]],
            "skipped": [_brief(r) for r in sel if r["price"] <= 0],
        }
    return out


def note_text(kind: str, day: date, loc_code: str, count_id: int) -> str:
    what = "илүүдэл" if kind == "income" else "дутагдал"
    name = location_name(loc_code)
    return f"{day.isoformat()}-ний тооллогын {what} · {loc_code}{' ' + name if name else ''}" + (f" (#{count_id})" if count_id else "")


def build_xlsx(rows: list[dict], *, kind: str, day: date, loc_code: str, count_id: int, customer: str,
               related_account: str, account: str, source_file: str = "") -> tuple[bytes, str, dict]:
    """Эрхэтийн импорт Excel. kind: income (inv_income 23 багана, огноо A) | expense (inv_csale 25 багана, огноо B).
    Нэг баримт: *-гүй (баримтын) талбар зөвхөн эхний мөрөнд, *-тэй (барааны) талбар мөр бүрт."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from app.services.erkhet_xlsx import apply_date_format, finalize_erkhet_xlsx

    sel = [r for r in rows if r["kind"] == kind]
    picked = [r for r in sel if r["price"] > 0]
    skipped = [r for r in sel if r["price"] <= 0]
    if not picked:
        raise ValueError(f"{'Илүүдэл' if kind == 'income' else 'Дутагдал'} бүхий (үнэтэй) мөр алга — файл гаргах зүйлгүй.")
    note = note_text(kind, day, loc_code, count_id)
    d = datetime.combine(day, datetime.min.time())
    wb = Workbook()
    ws = wb.active
    ws.title = "Import"
    if kind == "income":
        ws.append(INCOME_HEADERS)
        for i, r in enumerate(picked):
            f = i == 0
            ws.append([d if f else None, None, note if f else None, customer if f else None,
                       related_account if f else None, None, 0 if f else None, None, None, 0 if f else None,
                       0 if f else None, None, 0 if f else None, account, r["code"], loc_code,
                       _num(r["qty"]), _num(r["price"]), 1, _num(round(r["qty"] * r["price"], 2)), 0, 0, None])
        date_col = "A"
    else:
        ws.append(CSALE_HEADERS)
        for i, r in enumerate(picked):
            f = i == 0
            z = 0 if f else None                    # баримтын тийм/үгүй, дүн — зөвхөн эхний мөрөнд
            total = _num(round(r["qty"] * r["price"], 2))
            ws.append([None, d if f else None, note if f else None, customer if f else None,
                       related_account if f else None, None, None,      # ялгаатай харилцагч, борлуулагч
                       z, z, z, None, z, z, z,                         # НӨАТ-тэй, авто, дүн, үзүүлэлт · НХАТ-тай, авто, дүн
                       r["code"], account, loc_code, _num(r["qty"]), _num(r["price"]), total,
                       0, 0, 0, total, None])                          # мөр-НӨАТ, мөр-НХАТ, НХАТ үзүүлэлт, хөнгөлөгдсөн дүн, хариуцагч РД
        date_col = "B"
    apply_date_format(ws, date_col, 2, 2)

    if kind == "income":
        # Тайлбар хуудас — Эрхэт зөвхөн эхний «Import»-ийг уншдаг. Зарлагын (огноо B) файлд нэмэхгүй:
        # Excel-ээр дахин хадгалахад бүх хуудасны B баганыг огноо болгодог.
        ws2 = wb.create_sheet("Тайлбар")
        for col, w in zip("ABC", (34, 30, 70)):
            ws2.column_dimensions[col].width = w
        info = [
            ("Тооллого", f"#{count_id} · {source_file}", ""),
            ("Төрөл", KIND_LABEL[kind], "Тооллогоор илүү гарсан бараа (Тооллого − Програм > 0)."),
            ("Огноо · байршил", f"{day.isoformat()} · {loc_code} {location_name(loc_code)}", ""),
            ("Харилцагч · харьцсан данс · данс", f"{customer} · {related_account} · {account}", ""),
            ("Мөр · ширхэг · дүн", f"{len(picked)} · {_num(sum(r['qty'] for r in picked))} · "
                                   f"{round(sum(r['qty'] * r['price'] for r in picked), 2):,.2f}", ""),
            *[(f"Нэгж үнэ: {lab}", sum(1 for r in picked if r["price_src"] == s), "")
              for s, lab in PRICE_SRC_LABEL.items() if any(r["price_src"] == s for r in picked)],
            ("Үнэгүй тул ХАСАГДСАН", len(skipped),
             "0 үнээр орлого авбал жигнэсэн дундаж өртөг эвдэрнэ — доор жагсаав, гараар оруулна уу."),
            ("⚠ АНХААРУУЛГА", "Нэг л удаа импортлоно", "Дахин импортловол илүүдэл 2 удаа орлого болно."),
        ]
        for ri, (k, v, n) in enumerate(info, 1):
            ws2.cell(ri, 1, k).font = Font(bold=True)
            ws2.cell(ri, 2, v)
            ws2.cell(ri, 3, n).font = Font(size=9, color="7F8C8D")
        for c in (1, 2, 3):
            ws2.cell(len(info), c).fill = PatternFill("solid", fgColor="FDEBD0")
        ri = len(info) + 1
        for title, items in (("ШАЛГАХ: өртөг олдоогүй тул ЗАРАХ ҮНЭЭР бодсон", [r for r in picked if r["review"]]),
                             ("ХАСАГДСАН (үнэгүй) мөрүүд", skipped)):
            if not items:
                continue
            ri += 1
            ws2.cell(ri, 1, title).font = Font(bold=True, color="C0392B")
            for r in items:
                ri += 1
                ws2.cell(ri, 1, r["code"])
                ws2.cell(ri, 2, r["name"])
                ws2.cell(ri, 3, f"{_num(r['qty'])} ш × {r['price']:,.2f}")

    buf = io.BytesIO()
    wb.save(buf)
    data = finalize_erkhet_xlsx(buf.getvalue(), date_col=date_col)
    fname = f"{day:%Y%m%d}_toollogo{count_id or ''}_{loc_code}_{KIND_SLUG[kind]}.xlsx"
    stats = {"rows": len(picked), "skipped": len(skipped), "pieces": _num(sum(r["qty"] for r in picked)),
             "amount": round(sum(r["qty"] * r["price"] for r in picked), 2)}
    return data, fname, stats
