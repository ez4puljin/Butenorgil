"""Нэмэлт талбар (мастер) — үйлчилгээ.

Эх сурвалж (одоогийн жагсаалт):
  product  → app/data/outputs/master_latest.xlsx  («Нэгтгэл» sheet: Код, Нэр, Баркод …)
  customer → app/data/outputs/customer_info_last.xlsx (Код, Нэр, Бүлэг нэр, Утас, Банк дахь данс …)
Импорт бүрийн дараа `sync_after_import()` дуудагдана:
  1. Файлд байгаа код → бичлэг active, нэр/зангуу шинэчлэгдэнэ.
  2. Файлд байхгүй болсон код → зангуугаар (баркод, нэр / нэр, утас, данс) файлын
     ШИНЭ (бичлэггүй) кодуудаас ганц таарах олдвол автоматаар тэр код руу шилжинэ.
  3. Олдохгүй бол status="orphan" — «Холбоос» цэснээс гараар холбоно. Утга устахгүй.
Мөн `apply_to_excel()` нь мастер Excel-д идэвхтэй талбаруудыг багана болгон нэмнэ.
"""
from __future__ import annotations

import io
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

from sqlalchemy.orm import Session

from app.models.custom_master import ENTITIES, FIELD_TYPES, CustomField, CustomRecord

MASTER_PATH = Path("app/data/outputs/master_latest.xlsx")
CUSTOMER_PATH = Path("app/data/outputs/customer_info_last.xlsx")

ENTITY_LABEL = {"product": "Бараа материал", "customer": "Харилцагч"}

# Анх удаа (хүснэгт хоосон үед) үүсгэх талбарууд — дараа нь хэрэглэгч чөлөөтэй засна
DEFAULT_FIELDS: list[dict] = [
    {"entity": "product",  "key": "freight_class", "label": "Ачааны ангилал", "ftype": "select",
     "options": ["Хүнд", "Хөнгөн", "Цул"], "sort_order": 1},
    {"entity": "customer", "key": "sys_group", "label": "Бүлэг (систем)", "ftype": "select",
     "options": ["Нийлүүлэгч"], "sort_order": 1},
    {"entity": "customer", "key": "location", "label": "Байршил", "ftype": "text", "group_filter": "Нийлүүлэгч", "sort_order": 2},
    {"entity": "customer", "key": "settlement", "label": "Тооцоо", "ftype": "text", "group_filter": "Нийлүүлэгч", "sort_order": 3},
    {"entity": "customer", "key": "vat", "label": "НӨАТ", "ftype": "select",
     "options": ["НӨАТ-тэй", "НӨАТ-гүй"], "group_filter": "Нийлүүлэгч", "sort_order": 4},
]
GROUP_KEY = "sys_group"   # харилцагчийн «Бүлэг (систем)» талбарын түлхүүр (group_filter үүнтэй тулгана)


class CustomMasterError(Exception):
    pass


# ── Хэвийн болгох ────────────────────────────────────────────────────────────
def norm_text(s) -> str:
    if s is None:
        return ""
    s = str(s).strip()
    if s.lower() in ("nan", "none", ""):
        return ""
    return re.sub(r"\s+", " ", s).lower()


def norm_code(s) -> str:
    if s is None:
        return ""
    s = str(s).strip()
    if s.lower() in ("nan", "none"):
        return ""
    if s.endswith(".0"):
        s = s[:-2]
    return s


def norm_digits(s) -> str:
    return re.sub(r"\D", "", str(s or "")) if str(s or "").lower() not in ("nan", "none") else ""


def split_barcodes(s) -> list[str]:
    raw = str(s or "").strip()
    if raw.lower() in ("", "nan", "none"):
        return []
    out: list[str] = []
    for p in re.split(r"[\s,;|/]+", raw):
        p = p.strip()
        if not p or p.lower() == "nan":
            continue
        try:
            if "e" in p.lower() or "." in p:
                p = str(int(float(p)))
        except (TypeError, ValueError):
            pass
        if p not in out:
            out.append(p)
    return out


# ── Одоогийн жагсаалт (файлаас, mtime кэштэй) ────────────────────────────────
_rows_cache: dict[str, tuple[float, str, list[dict]]] = {}


def _read_excel(path: Path, **kw):
    import pandas as pd
    try:
        return pd.read_excel(str(path), engine="calamine", **kw)
    except Exception:
        return pd.read_excel(str(path), **kw)


def _cell(row, col: str) -> str:
    v = row.get(col)
    if v is None:
        return ""
    try:
        import pandas as pd
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    s = str(v).strip()
    return "" if s.lower() in ("nan", "none") else s


def _load_product_rows(path: Path) -> list[dict]:
    df = _read_excel(path, sheet_name=0, dtype=str)
    if "Код" not in df.columns:
        raise CustomMasterError("Мастер файлд «Код» багана алга")
    rows: list[dict] = []
    seen: set[str] = set()
    for _, r in df.iterrows():
        code = norm_code(_cell(r, "Код"))
        if not code or code in seen:
            continue
        seen.add(code)
        name = _cell(r, "Нэр")
        rows.append({
            "code": code, "name": name,
            "info": {"category": _cell(r, "Ангилал нэр"), "brand": _cell(r, "Брэнд нэр"),
                     "barcode": _cell(r, "Баркод"), "location_tag": _cell(r, "Байршил tag")},
            "anchors": {"name": norm_text(name), "barcodes": split_barcodes(_cell(r, "Баркод"))},
        })
    return rows


def _load_customer_rows(path: Path) -> list[dict]:
    df = _read_excel(path, sheet_name=0, dtype=str)
    if "Код" not in df.columns:
        raise CustomMasterError("Харилцагчийн файлд «Код» багана алга")
    rows: list[dict] = []
    seen: set[str] = set()
    for _, r in df.iterrows():
        code = norm_code(_cell(r, "Код"))
        if not code or code in seen:
            continue
        seen.add(code)
        name = _cell(r, "Нэр")
        phone = norm_digits(_cell(r, "Утас"))
        acct = norm_digits(_cell(r, "Банк дахь данс"))
        rows.append({
            "code": code, "name": name,
            "info": {"group_code": _cell(r, "Бүлэг код"), "group_name": _cell(r, "Бүлэг нэр"),
                     "email": _cell(r, "мэйл"), "address": _cell(r, "Хаяг"), "phone": _cell(r, "Утас"),
                     "vat_type": _cell(r, "НӨАТ төрөл"), "bank": _cell(r, "Банк"), "bank_acct": _cell(r, "Банк дахь данс")},
            "anchors": {"name": norm_text(name), "phone": phone if len(phone) >= 6 else "",
                        "bank_acct": acct if len(acct) >= 6 else "", "email": norm_text(_cell(r, "мэйл"))},
        })
    return rows


def source_path(entity: str) -> Path:
    return MASTER_PATH if entity == "product" else CUSTOMER_PATH


def current_rows(entity: str, path: Optional[Path] = None) -> list[dict]:
    """Файлаас одоогийн жагсаалт (mtime өөрчлөгдөөгүй бол кэшээс)."""
    if entity not in ENTITIES:
        raise CustomMasterError(f"Буруу entity: {entity}")
    p = Path(path) if path else source_path(entity)
    if not p.exists():
        return []
    mtime = p.stat().st_mtime
    c = _rows_cache.get(entity)
    if c and c[0] == mtime and c[1] == str(p):
        return c[2]
    rows = _load_product_rows(p) if entity == "product" else _load_customer_rows(p)
    _rows_cache[entity] = (mtime, str(p), rows)
    return rows


def invalidate_rows(entity: str | None = None):
    if entity:
        _rows_cache.pop(entity, None)
    else:
        _rows_cache.clear()


# ── Талбар ───────────────────────────────────────────────────────────────────
def field_options(f: CustomField) -> list[str]:
    try:
        v = json.loads(f.options or "[]")
        return [str(x) for x in v] if isinstance(v, list) else []
    except Exception:
        return []


def field_to_dict(f: CustomField) -> dict:
    return {"id": f.id, "entity": f.entity, "key": f.key, "label": f.label, "ftype": f.ftype,
            "options": field_options(f), "group_filter": f.group_filter or "", "sort_order": f.sort_order,
            "is_active": bool(f.is_active)}


def active_fields(db: Session, entity: str) -> list[CustomField]:
    return (db.query(CustomField).filter(CustomField.entity == entity, CustomField.is_active == True)  # noqa: E712
            .order_by(CustomField.sort_order.asc(), CustomField.id.asc()).all())


def validate_values(fields: Iterable[CustomField], values: dict) -> dict:
    """Утгуудыг талбарын төрлөөр шалгаж цэвэрлэнэ. Хоосон утга → түлхүүр хасагдана."""
    fmap = {f.key: f for f in fields}
    out: dict = {}
    for k, v in (values or {}).items():
        f = fmap.get(k)
        if f is None:
            raise CustomMasterError(f"Тодорхойгүй талбар: {k}")
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        if f.ftype == "number":
            try:
                num = float(str(v).replace(",", "."))
            except ValueError:
                raise CustomMasterError(f"«{f.label}» тоо байх ёстой")
            out[k] = int(num) if num.is_integer() else num
        elif f.ftype == "bool":
            out[k] = bool(v) if not isinstance(v, str) else v.strip().lower() in ("1", "true", "тийм", "yes", "да")
        elif f.ftype == "select":
            s = str(v).strip()
            opts = field_options(f)
            if opts and s not in opts:
                raise CustomMasterError(f"«{f.label}»: «{s}» сонголт байхгүй ({', '.join(opts)})")
            out[k] = s
        elif f.ftype == "date":
            s = str(v).strip()[:10]
            try:
                datetime.strptime(s, "%Y-%m-%d")
            except ValueError:
                raise CustomMasterError(f"«{f.label}» огноо YYYY-MM-DD байх ёстой")
            out[k] = s
        else:
            out[k] = str(v).strip()[:500]
    return out


def ensure_defaults(db: Session):
    """Талбарын хүснэгт хоосон бол анхны талбаруудыг үүсгэнэ (нэг л удаа)."""
    if db.query(CustomField.id).first() is not None:
        return
    for i, d in enumerate(DEFAULT_FIELDS):
        db.add(CustomField(entity=d["entity"], key=d["key"], label=d["label"], ftype=d.get("ftype", "text"),
                           options=json.dumps(d.get("options", []), ensure_ascii=False),
                           group_filter=d.get("group_filter", ""), sort_order=d.get("sort_order", i), is_active=True))
    db.commit()


# ── Бичлэг ───────────────────────────────────────────────────────────────────
def _j(s: str, default):
    try:
        v = json.loads(s or "")
        return v if v is not None else default
    except Exception:
        return default


def record_to_dict(r: CustomRecord) -> dict:
    return {"code": r.code, "name": r.name, "values": _j(r.values, {}), "status": r.status,
            "anchors": _j(r.anchors, {}), "prev_codes": _j(r.prev_codes, []),
            "updated_by": r.updated_by, "updated_at": r.updated_at.isoformat() if r.updated_at else None,
            "seen_at": r.seen_at.isoformat() if r.seen_at else None}


def get_record(db: Session, entity: str, code: str) -> Optional[CustomRecord]:
    return db.query(CustomRecord).filter(CustomRecord.entity == entity, CustomRecord.code == code).first()


def values_map(db: Session, entity: str, codes: Iterable[str] | None = None) -> dict[str, dict]:
    """{code: values} — системийн бусад хэсэгт (хайлт, тайлан) ашиглана."""
    q = db.query(CustomRecord).filter(CustomRecord.entity == entity)
    if codes is not None:
        codes = [c for c in set(codes) if c]
        if not codes:
            return {}
        q = q.filter(CustomRecord.code.in_(codes))
    return {r.code: _j(r.values, {}) for r in q.all()}


def upsert_record(db: Session, entity: str, code: str, values: dict, by: str,
                  rows: list[dict] | None = None) -> CustomRecord:
    fields = active_fields(db, entity)
    clean = validate_values(fields, values)
    rows = rows if rows is not None else current_rows(entity)
    row = next((x for x in rows if x["code"] == code), None)
    r = get_record(db, entity, code)
    if r is None:
        if row is None:
            raise CustomMasterError(f"«{code}» код одоогийн {ENTITY_LABEL[entity].lower()}ийн жагсаалтад алга")
        r = CustomRecord(entity=entity, code=code, status="active")
        db.add(r)
    if row is not None:
        r.name = row["name"]
        r.anchors = json.dumps(row["anchors"], ensure_ascii=False)
        r.status = "active"
        r.seen_at = datetime.utcnow()
    r.values = json.dumps(clean, ensure_ascii=False)
    r.updated_by = by
    r.updated_at = datetime.utcnow()
    db.commit()
    return r


# ── Зангуугаар тулгах ────────────────────────────────────────────────────────
STRONG = {"product": ("barcodes", "name"), "customer": ("name", "phone", "bank_acct")}
ANCHOR_LABEL = {"barcodes": "баркод", "name": "нэр", "phone": "утас", "bank_acct": "данс", "email": "и-мэйл"}


def _anchor_index(entity: str, rows: list[dict]) -> dict[str, dict[str, list[str]]]:
    """{anchor_kind: {value: [codes]}}"""
    idx: dict[str, dict[str, list[str]]] = {k: {} for k in STRONG[entity]}
    for row in rows:
        a = row["anchors"]
        for k in STRONG[entity]:
            vals = a.get(k) or []
            if isinstance(vals, str):
                vals = [vals] if vals else []
            for v in vals:
                idx[k].setdefault(v, []).append(row["code"])
    return idx


def _match_candidates(entity: str, anchors: dict, idx: dict, allowed: set[str]) -> dict[str, list[str]]:
    """{code: [anchor_kind,...]} — зангуу таарсан кодууд (allowed доторх)."""
    found: dict[str, list[str]] = {}
    for k in STRONG[entity]:
        vals = anchors.get(k) or []
        if isinstance(vals, str):
            vals = [vals] if vals else []
        for v in vals:
            for c in idx[k].get(v, []):
                if c in allowed:
                    found.setdefault(c, [])
                    if k not in found[c]:
                        found[c].append(k)
    return found


def suggestions(db: Session, entity: str, rec: CustomRecord, rows: list[dict] | None = None, limit: int = 5) -> list[dict]:
    """Гараар холбоход санал болгох шинэ кодууд (оноогоор эрэмбэлсэн)."""
    rows = rows if rows is not None else current_rows(entity)
    recorded = {r.code for r in db.query(CustomRecord.code).filter(CustomRecord.entity == entity).all()}
    allowed = {x["code"] for x in rows if x["code"] not in recorded or x["code"] == rec.code}
    idx = _anchor_index(entity, rows)
    found = _match_candidates(entity, _j(rec.anchors, {}), idx, allowed)
    by_code = {x["code"]: x for x in rows}
    out = []
    weight = {"barcodes": 3, "bank_acct": 3, "phone": 2, "name": 2}
    for c, kinds in found.items():
        out.append({"code": c, "name": by_code[c]["name"], "info": by_code[c]["info"],
                    "matched": [ANCHOR_LABEL[k] for k in kinds], "score": sum(weight.get(k, 1) for k in kinds)})
    # Нэрний эхний үгээр сул санал (зангуу огт таараагүй үед)
    if not out:
        first = (norm_text(rec.name).split(" ") or [""])[0]
        if len(first) >= 3:
            for x in rows:
                if x["code"] in allowed and x["anchors"].get("name", "").startswith(first):
                    out.append({"code": x["code"], "name": x["name"], "info": x["info"], "matched": ["нэрний эхлэл"], "score": 1})
                    if len(out) >= limit:
                        break
    out.sort(key=lambda d: (-d["score"], d["code"]))
    return out[:limit]


def sync_after_import(db: Session, entity: str, rows: list[dict] | None = None, by: str = "import") -> dict:
    """Импортын дараа: active/orphan төлөв, автомат холболт."""
    rows = rows if rows is not None else current_rows(entity)
    if not rows:
        return {"entity": entity, "rows": 0, "skipped": True}
    by_code = {x["code"]: x for x in rows}
    recs = db.query(CustomRecord).filter(CustomRecord.entity == entity).all()
    rec_codes = {r.code for r in recs}
    now = datetime.utcnow()

    orphans: list[CustomRecord] = []
    for r in recs:
        row = by_code.get(r.code)
        if row is not None:
            r.name = row["name"]
            r.anchors = json.dumps(row["anchors"], ensure_ascii=False)
            r.status = "active"
            r.seen_at = now
        else:
            orphans.append(r)

    relinked: list[dict] = []
    if orphans:
        allowed = {c for c in by_code if c not in rec_codes}     # бичлэггүй шинэ кодууд
        idx = _anchor_index(entity, rows)
        claims: dict[str, list[tuple[CustomRecord, list[str]]]] = {}
        for r in orphans:
            found = _match_candidates(entity, _j(r.anchors, {}), idx, allowed)
            if len(found) == 1:
                (c, kinds), = found.items()
                claims.setdefault(c, []).append((r, kinds))
        for c, lst in claims.items():
            if len(lst) != 1:
                continue      # нэг шинэ кодыг хоёр хуучин бичлэг нэхэж байвал гараар
            r, kinds = lst[0]
            hist = _j(r.prev_codes, [])
            hist.append({"code": r.code, "name": r.name, "at": now.isoformat(timespec="seconds"), "by": by,
                         "matched": [ANCHOR_LABEL[k] for k in kinds]})
            row = by_code[c]
            relinked.append({"old": r.code, "new": c, "name": row["name"], "matched": [ANCHOR_LABEL[k] for k in kinds]})
            r.code = c
            r.name = row["name"]
            r.anchors = json.dumps(row["anchors"], ensure_ascii=False)
            r.prev_codes = json.dumps(hist, ensure_ascii=False)
            r.status = "active"
            r.seen_at = now
            orphans.remove(r)
        for r in orphans:
            r.status = "orphan"
    db.commit()
    return {"entity": entity, "rows": len(rows), "records": len(recs), "relinked": relinked, "orphans": len(orphans)}


def relink(db: Session, entity: str, old_code: str, new_code: str, by: str, rows: list[dict] | None = None) -> CustomRecord:
    """Гараар холбох: хуучин бичлэгийн утгыг шинэ код руу шилжүүлнэ (шинэ кодод бичлэг байвал хоосон талбарыг нөхнө)."""
    rows = rows if rows is not None else current_rows(entity)
    row = next((x for x in rows if x["code"] == new_code), None)
    if row is None:
        raise CustomMasterError(f"«{new_code}» код одоогийн жагсаалтад алга")
    old = get_record(db, entity, old_code)
    if old is None:
        raise CustomMasterError(f"«{old_code}» бичлэг олдсонгүй")
    if old_code == new_code:
        return old
    now = datetime.utcnow()
    hist = _j(old.prev_codes, [])
    hist.append({"code": old.code, "name": old.name, "at": now.isoformat(timespec="seconds"), "by": by, "matched": ["гараар"]})
    target = get_record(db, entity, new_code)
    if target is not None:
        merged = {**_j(old.values, {}), **_j(target.values, {})}      # шинэ дээрх утга давуу, хоосон бол хуучнаас
        target.values = json.dumps(merged, ensure_ascii=False)
        target.prev_codes = json.dumps(_j(target.prev_codes, []) + hist, ensure_ascii=False)
        target.updated_by = by
        target.updated_at = now
        db.delete(old)
        rec = target
    else:
        old.code = new_code
        old.prev_codes = json.dumps(hist, ensure_ascii=False)
        old.updated_by = by
        old.updated_at = now
        rec = old
    rec.name = row["name"]
    rec.anchors = json.dumps(row["anchors"], ensure_ascii=False)
    rec.status = "active"
    rec.seen_at = now
    db.commit()
    return rec


# ── Excel ────────────────────────────────────────────────────────────────────
def _excel_value(f: CustomField, v):
    if v is None:
        return None
    if f.ftype == "bool":
        return "Тийм" if v else "Үгүй"
    return v


def apply_to_excel(db: Session, entity: str, path: Optional[Path] = None, code_header: str = "Код") -> dict:
    """Мастер Excel-ийн эхний sheet-д идэвхтэй талбаруудыг багана болгон нэмнэ/шинэчилнэ.
    Байгаа багануудыг хөдөлгөхгүй (бусад parser байрлал/нэрээр уншдаг) — зөвхөн төгсгөлд нэмнэ."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    p = Path(path) if path else source_path(entity)
    if not p.exists():
        return {"skipped": "file missing"}
    fields = active_fields(db, entity)
    if not fields:
        return {"skipped": "no fields"}
    vals = values_map(db, entity)
    wb = load_workbook(str(p))
    ws = wb.worksheets[0]
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    code_col = headers.get(code_header)
    if not code_col:
        return {"skipped": f"'{code_header}' header not found"}
    col_of: dict[str, int] = {}
    next_col = ws.max_column + 1
    for f in fields:
        if f.label in headers:
            col_of[f.key] = headers[f.label]
        else:
            col_of[f.key] = next_col
            c = ws.cell(row=1, column=next_col, value=f.label)
            c.font = Font(bold=True)
            next_col += 1
    written = 0
    for row in range(2, ws.max_row + 1):
        code = norm_code(ws.cell(row=row, column=code_col).value)
        v = vals.get(code)
        for f in fields:
            cell = ws.cell(row=row, column=col_of[f.key])
            cell.value = _excel_value(f, (v or {}).get(f.key)) if v else None
        if v:
            written += 1
    wb.save(str(p))
    invalidate_rows(entity)
    return {"columns": [f.label for f in fields], "rows_with_values": written}


def build_export(db: Session, entity: str) -> bytes:
    """Одоогийн жагсаалт + нэмэлт баганууд (+ талбарууд, холбоос хүлээж буй) Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rows = current_rows(entity)
    fields = active_fields(db, entity)
    recs = {r.code: r for r in db.query(CustomRecord).filter(CustomRecord.entity == entity).all()}
    if entity == "product":
        base_cols = [("code", "Код"), ("name", "Нэр"), ("category", "Ангилал нэр"), ("brand", "Брэнд нэр"),
                     ("barcode", "Баркод"), ("location_tag", "Байршил tag")]
    else:
        base_cols = [("code", "Код"), ("name", "Нэр"), ("group_code", "Бүлэг код"), ("group_name", "Бүлэг нэр"),
                     ("email", "мэйл"), ("address", "Хаяг"), ("phone", "Утас"), ("vat_type", "НӨАТ төрөл"),
                     ("bank", "Банк"), ("bank_acct", "Банк дахь данс")]
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", start_color="DDEBF7")
    custom_fill = PatternFill("solid", start_color="FFF2CC")

    wb = Workbook()
    ws = wb.active
    ws.title = ENTITY_LABEL[entity]
    heads = [h for _, h in base_cols] + [f.label for f in fields]
    ws.append(heads)
    for i, c in enumerate(ws[1], start=1):
        c.font = Font(bold=True)
        c.fill = custom_fill if i > len(base_cols) else head_fill
        c.border = border
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        rec = recs.get(row["code"])
        vals = _j(rec.values, {}) if rec else {}
        line = [row["code"], row["name"]] + [row["info"].get(k, "") for k, _ in base_cols[2:]]
        line += [_excel_value(f, vals.get(f.key)) for f in fields]
        ws.append(line)
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in r:
            c.border = border
    widths = {"Код": 10, "Нэр": 42, "Ангилал нэр": 22, "Брэнд нэр": 22, "Баркод": 18, "Байршил tag": 20,
              "Хаяг": 28, "мэйл": 20, "Утас": 14, "Банк дахь данс": 18, "Бүлэг нэр": 22}
    for i, h in enumerate(heads, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 16)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Талбарууд")
    ws2.append(["Түлхүүр", "Нэр", "Төрөл", "Сонголт", "Зөвхөн бүлэг", "Дараалал"])
    for f in fields:
        ws2.append([f.key, f.label, f.ftype, ", ".join(field_options(f)), f.group_filter or "", f.sort_order])
    for c in ws2[1]:
        c.font = Font(bold=True)

    orphans = [r for r in recs.values() if r.status == "orphan"]
    if orphans:
        ws3 = wb.create_sheet("Холбоос хүлээж буй")
        ws3.append(["Хуучин код", "Нэр", "Утгууд", "Сүүлд файлд байсан"])
        for r in orphans:
            vals = _j(r.values, {})
            ws3.append([r.code, r.name, "; ".join(f"{k}={v}" for k, v in vals.items()),
                        r.seen_at.strftime("%Y-%m-%d") if r.seen_at else ""])
        for c in ws3[1]:
            c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def after_import(db: Session, entity: str, path: Optional[Path] = None, by: str = "import") -> dict:
    """Импортын төгсгөлд дуудна: холболт + Excel багана. Алдаа гарвал импортыг унагахгүй."""
    out: dict = {}
    try:
        invalidate_rows(entity)
        rows = current_rows(entity, path)
        out["sync"] = sync_after_import(db, entity, rows, by=by)
    except Exception as e:   # noqa: BLE001
        db.rollback()
        out["sync_error"] = f"{type(e).__name__}: {e}"
    try:
        out["excel"] = apply_to_excel(db, entity, path)
        current_rows(entity, path)   # кэшийг урьдчилан дулаацуулна (дараагийн хүсэлт хүлээхгүй)
    except Exception as e:   # noqa: BLE001
        out["excel_error"] = f"{type(e).__name__}: {e}"
    return out


def bulk_update(db: Session, entity: str, codes: list[str], patch: dict, by: str,
                rows: list[dict] | None = None) -> dict:
    """Олон бичлэгт нэг дор утга оноох. patch: {key: value}; value хоосон/None бол тэр талбарыг хоослоно.
    «Зөвхөн бүлэгт» талбарыг тухайн бүлгийн бус бичлэгт оноохгүй (skipped_group)."""
    fields = active_fields(db, entity)
    fmap = {f.key: f for f in fields}
    for k in patch:
        if k not in fmap:
            raise CustomMasterError(f"Тодорхойгүй талбар: {k}")
    clear_keys = [k for k, v in patch.items() if v is None or (isinstance(v, str) and not v.strip())]
    clean = validate_values(fields, {k: v for k, v in patch.items() if k not in clear_keys})
    if not clean and not clear_keys:
        raise CustomMasterError("Оноох утга алга")
    rows = rows if rows is not None else current_rows(entity)
    by_code = {x["code"]: x for x in rows}
    codes = [c for c in dict.fromkeys(codes) if c]
    recs = {r.code: r for r in db.query(CustomRecord).filter(CustomRecord.entity == entity, CustomRecord.code.in_(codes)).all()} if codes else {}
    now = datetime.utcnow()
    updated = 0
    skipped_unknown: list[str] = []
    skipped_group = 0
    unchanged = 0
    for code in codes:
        row = by_code.get(code)
        rec = recs.get(code)
        if row is None and rec is None:
            skipped_unknown.append(code)
            continue
        cur = _j(rec.values, {}) if rec else {}
        new = dict(cur)
        eff_group = str(clean.get(GROUP_KEY, cur.get(GROUP_KEY, "")))
        skipped_here = False
        for k, v in clean.items():
            f = fmap[k]
            if f.group_filter and eff_group != f.group_filter:
                skipped_here = True
                continue
            new[k] = v
        for k in clear_keys:
            new.pop(k, None)
        if skipped_here:
            skipped_group += 1
        if new == cur:
            if not skipped_here:
                unchanged += 1
            continue
        if rec is None:
            rec = CustomRecord(entity=entity, code=code, status="active")
            db.add(rec)
        if row is not None:
            rec.name = row["name"]
            rec.anchors = json.dumps(row["anchors"], ensure_ascii=False)
            rec.status = "active"
            rec.seen_at = now
        rec.values = json.dumps(new, ensure_ascii=False)
        rec.updated_by = by
        rec.updated_at = now
        updated += 1
    db.commit()
    return {"updated": updated, "selected": len(codes), "unchanged": unchanged, "skipped_group": skipped_group, "skipped_unknown": skipped_unknown[:20]}
